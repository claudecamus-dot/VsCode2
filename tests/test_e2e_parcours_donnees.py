"""Parcours de DONNÉES dans un vrai navigateur (comité e2e du 2026-10-14,
dernière section du triage) : quatre parcours que les e2e voisins ne
jouaient pas, chacun par des clics RÉELS, avec une assertion sur ce que
l'utilisateur VOIT (`nav.texte()` = innerText) et la persistance relue
(rechargement de la page ET lecture de la base).

1. recommandations : « Générer (IA) » (chemin réel sous la coupe Ollama :
   message d'échec lisible, rien d'écrasé) puis édition d'un champ d'axe et de
   recommandation (autosave htmx), rechargement ;
2. sauvegarde du brouillon d'enregistrement : le micro est un faux
   périphérique Chromium, « Démarrer » puis « Arrêter » envoient l'audio brut à
   `/interviews/record/backup` ; l'onglet Backup de la mission le liste, puis
   « Supprimer » l'efface (fichier ET écran) ;
3. import d'un entretien existant (.docx) : upload multipart puis revue puis
   « Valider l'import » (deux temps). Sous la coupe, l'extraction IA échoue
   lisiblement (chemin réel) ; le chemin heureux passe par un serveur dont
   SEULE l'extraction IA est remplacée (`_serveur_extraction_factice`) ;
4. trames : créer (thème + question), éditer (autosave), importer un .docx
   (analyse déterministe, sans IA) puis revue puis validation, supprimer.

Règles tenues ici : clics à la souris (visibles ET non recouverts), polling
plutôt que `time.sleep` fixe, aucune erreur réseau/JS à chaque étape, une
mission semée par test, base et enregistrements dans le dossier temporaire du
module. Aucun test ne dépend d'une IA réelle.

Même garde que les voisins : sans navigateur Chromium le module est sauté,
sauf `E2E_OBLIGATOIRE=1` (alors c'est une ERREUR).
"""
from __future__ import annotations

import contextlib
import json
import os
import shutil
import subprocess
import sys
import tempfile
import time
import urllib.request
from pathlib import Path

import pytest

_OBLIGATOIRE = bool(os.environ.get("E2E_OBLIGATOIRE"))
try:
    import websockets  # noqa: F401 — fourni par uvicorn[standard]
except ModuleNotFoundError:
    if _OBLIGATOIRE:
        raise RuntimeError(
            "E2E_OBLIGATOIRE est posé mais `websockets` (uvicorn[standard]) est absent"
        ) from None
    pytest.skip("websockets (uvicorn[standard]) requis", allow_module_level=True)

import navigateur_cdp  # noqa: E402
from e2e_serveur import _arreter, serveur_uvicorn  # noqa: E402
from navigateur_cdp import Navigateur, port_libre, trouver_navigateur  # noqa: E402

_NAVIGATEUR = trouver_navigateur()

if _NAVIGATEUR is None and _OBLIGATOIRE:
    raise RuntimeError(
        "E2E_OBLIGATOIRE est posé mais aucun navigateur Chromium (Edge/Chrome) n'est trouvé"
    )

pytestmark = pytest.mark.skipif(
    _NAVIGATEUR is None,
    reason="aucun navigateur Chromium (Edge/Chrome) trouvé — E2E_NAVIGATEUR pour en forcer un",
)

_RACINE = Path(__file__).resolve().parents[1]
_WAV = _RACINE / "tests" / "fixtures" / "audio" / "tts_anglais_synthetique.wav"


# --------------------------------------------------------------------------- #
# Plomberie locale
# --------------------------------------------------------------------------- #
@pytest.fixture(scope="module")
def serveur(tmp_path_factory: pytest.TempPathFactory):
    """Un vrai uvicorn, Ollama injoignable exprès (et coupé s'il l'est encore).
    Rend `(base_url, dossier)` ; la base est `dossier / "e2e.db"`."""
    dossier = tmp_path_factory.mktemp("e2e-donnees")
    with serveur_uvicorn(dossier, OLLAMA_HOST="http://127.0.0.1:9") as base:
        yield base, dossier


@pytest.fixture
def nav():
    profil = Path(tempfile.mkdtemp(prefix="e2e-nav-"))
    navigateur = None
    try:
        navigateur = Navigateur(_NAVIGATEUR, profil / "profil")
        yield navigateur
    finally:
        if navigateur is not None:
            navigateur.fermer()
        shutil.rmtree(profil, ignore_errors=True)


@pytest.fixture
def nav_micro(monkeypatch: pytest.MonkeyPatch):
    """Navigateur au faux micro (clip de parole synthétique rejoué en boucle)."""
    assert _WAV.exists(), _WAV
    monkeypatch.setattr(navigateur_cdp, "_DRAPEAUX_LANCEMENT", [
        *navigateur_cdp._DRAPEAUX_LANCEMENT,
        "--use-fake-device-for-media-stream",
        "--use-fake-ui-for-media-stream",
        f"--use-file-for-fake-audio-capture={_WAV}",
    ])
    profil = Path(tempfile.mkdtemp(prefix="e2e-mic-"))
    navigateur = None
    try:
        navigateur = Navigateur(_NAVIGATEUR, profil / "profil")
        yield navigateur
    finally:
        if navigateur is not None:
            navigateur.fermer()
        shutil.rmtree(profil, ignore_errors=True)


def _python_sur_base(base_db: Path, code: str, *args: str) -> str:
    env = dict(os.environ, APP_DB_PATH=str(base_db), PYTHONUTF8="1")
    res = subprocess.run([sys.executable, "-c", code, *args], cwd=_RACINE, env=env,
                         capture_output=True, text=True, timeout=60, encoding="utf-8")
    assert res.returncode == 0, res.stderr
    return res.stdout.strip().splitlines()[-1]


def _sans_erreur(nav: Navigateur, etape: str, *, sauf: tuple[str, ...] = ()) -> None:
    """Aucune réponse 4xx/5xx, ressource perdue, exception JS ou erreur CDP
    (`sauf` : fins d'URL dont l'erreur est ATTENDUE et déjà asserée)."""
    nav.drainer()
    inattendues = [r for r in nav.erreurs()
                   if not any(str(r.get("url", "")).endswith(s) for s in sauf)]
    assert not inattendues, f"{etape} : réponses en erreur {inattendues}"
    assert not nav.exceptions, f"{etape} : exceptions JS {nav.exceptions}"
    assert not nav.erreurs_cdp, f"{etape} : erreur CDP {nav.erreurs_cdp}"
    assert "Origine non autorisée" not in nav.texte(), f"{etape} : 403 CSRF rendu à l'écran"


def _attendre(condition, message: str, delai_s: float = 10.0, pas_s: float = 0.2):
    fin = time.monotonic() + delai_s
    while time.monotonic() < fin:
        valeur = condition()
        if valeur:
            return valeur
        time.sleep(pas_s)
    raise AssertionError(message)


def _attendre_visible(nav: Navigateur, selecteur: str, attendu: str, delai_s: float = 10.0) -> None:
    """`attendu` dans le texte VISIBLE (innerText, insensible à la casse :
    `text-transform`) de l'élément."""
    _attendre(
        lambda: nav.evaluer(
            "(function(){var e=document.querySelector(" + json.dumps(selecteur) + ");"
            " return !!(e && e.innerText.toLowerCase().indexOf(" + json.dumps(attendu.lower()) + ")>=0);})()"
        ),
        f"« {attendu} » jamais visible dans {selecteur}", delai_s,
    )


def _saisir_au_clavier(nav: Navigateur, selecteur: str, valeur: str) -> None:
    """Autosaves htmx à `keyup` : `nav.remplir` n'émet que input/change, on
    termine donc par un `keyup`, comme un utilisateur qui tape."""
    nav.remplir(selecteur, valeur)
    ok = nav.evaluer(
        "(function(){var e=document.querySelector(" + json.dumps(selecteur) + ");"
        " if(!e) return false;"
        " e.dispatchEvent(new KeyboardEvent('keyup',{bubbles:true,key:'a'}));"
        " return true;})()"
    )
    assert ok, f"champ introuvable : {selecteur}"


def _valeur(nav: Navigateur, selecteur: str) -> str:
    return nav.evaluer(
        "(function(){var e=document.querySelector(" + json.dumps(selecteur) + ");"
        " return e ? e.value : null;})()")


# Message d'échec IA : l'un de DEUX textes selon l'état de `_call_ollama_once`
# (coupe temporaire d'Ollama, levée le 2026-10-19, ou réseau). On exige LA variante qui
# correspond au code réel, pas « l'une ou l'autre ».
_MSG_COUPE = "Ollama est désactivé temporairement"
_MSG_RESEAU = "Impossible de joindre Ollama sur http://127.0.0.1:9"


_coupe_ollama_active = navigateur_cdp.coupe_ollama_active  # détection AST partagée


def _assert_echec_ia_rendu(nav: Navigateur) -> str:
    """UN bandeau `.ai-error` VISIBLE, « ⚠ » + le texte exact de la variante
    attendue du service IA. Rend le message."""
    visibles = nav.evaluer(
        "Array.from(document.querySelectorAll('.ai-error')).filter(function(e){"
        "return e.offsetParent !== null;}).map(function(e){return e.innerText;})")
    assert len(visibles) == 1, f"un seul bandeau .ai-error visible attendu : {visibles}"
    message = visibles[0].strip()
    assert message.startswith("⚠ "), message
    attendu = _MSG_COUPE if _coupe_ollama_active() else _MSG_RESEAU
    assert attendu in message, f"« {attendu} » attendu : {message[:300]}"
    return message


# --------------------------------------------------------------------------- #
# 1. Recommandations : générer (échec lisible) puis éditer un champ
# --------------------------------------------------------------------------- #
_SEED_RECO = r"""
from app.db import SessionLocal, init_db
from app.models import GlobalSynthesis, Interview, Mission, Question, Theme, Trame
init_db()
db = SessionLocal()
m = Mission(name="E2E reco")
db.add(m); db.flush()
tr = Trame(mission_id=m.id); db.add(tr); db.flush()
th = Theme(trame_id=tr.id, title="Gouvernance", position=0); db.add(th); db.flush()
db.add(Question(theme_id=th.id, label="Comment décidez-vous ?", position=0))
db.add(Interview(mission_id=m.id, interviewee_name="Témoin reco", status="done"))
db.add(GlobalSynthesis(mission_id=m.id, status="generated",
                       contexte="Un contexte de test.", points_amelioration="- Silos entre équipes"))
db.commit()
print(m.id)
"""

_SEED_AXE = r"""
import sys
from app.db import SessionLocal
from app.models import Recommendation, RecommendationAxis
db = SessionLocal()
a = RecommendationAxis(mission_id=int(sys.argv[1]), title="Axe initial", position=0)
db.add(a); db.flush()
r = Recommendation(axis_id=a.id, title="Reco initiale", position=0)
db.add(r); db.commit()
print(a.id, r.id)
"""

_LIRE_RECO = r"""
import sys
from app.db import SessionLocal
from app.models import Recommendation, RecommendationAxis
db = SessionLocal()
a = db.get(RecommendationAxis, int(sys.argv[1])); r = db.get(Recommendation, int(sys.argv[2]))
print(repr((a.title, a.edite, r.title, r.objectif, r.edite)))
"""

_COMPTER_AXES = r"""
import sys
from app.db import SessionLocal
from app.models import RecommendationAxis
print(db_count := SessionLocal().query(RecommendationAxis).filter_by(mission_id=int(sys.argv[1])).count())
"""


def test_recommandations_generer_sous_la_coupe_puis_editer_un_champ(serveur, nav: Navigateur) -> None:
    """« Générer (IA) » cliqué pour de vrai : l'IA est coupée, l'écran dit
    POURQUOI (bandeau d'échec lisible) et ne crée aucun axe. Puis, sur des axes
    présents, un titre d'axe et un champ de recommandation retapés au clavier
    s'enregistrent (« ✓ enregistré »), survivent au rechargement et sont lus en
    base avec leur drapeau « édité »."""
    base, dossier = serveur
    base_db = dossier / "e2e.db"
    mid = _python_sur_base(base_db, _SEED_RECO)

    nav.naviguer(f"{base}/missions/{mid}/recommandations")
    _sans_erreur(nav, "Ouvrir les recommandations")
    assert "Aucune recommandation pour l'instant." in nav.texte()
    bouton = f"form[action='/missions/{mid}/recommandations/generate'] button[type=submit]"
    assert nav.evaluer(
        "document.querySelector(" + json.dumps(bouton) + ").innerText.trim()") == "Générer (IA)"
    nav.cliquer_et_attendre(bouton)
    _sans_erreur(nav, "Générer les recommandations")
    _assert_echec_ia_rendu(nav)
    assert "Aucune recommandation pour l'instant." in nav.texte(), "l'échec a laissé des axes"
    assert _python_sur_base(base_db, _COMPTER_AXES, mid) == "0"

    aid, rid = _python_sur_base(base_db, _SEED_AXE, mid).split()
    nav.naviguer(f"{base}/missions/{mid}/recommandations")
    assert _valeur(nav, f"input[hx-post='/recommandations/axes/{aid}/field']") == "Axe initial"
    _saisir_au_clavier(nav, f"input[hx-post='/recommandations/axes/{aid}/field']", "Axe gouvernance partagée")
    _attendre_visible(nav, f"#axis-saved-{aid}", "enregistré")
    _saisir_au_clavier(
        nav, f"#reco-card-{rid} textarea[hx-post='/recommandations/{rid}/field'][hx-vals*='objectif']",
        "Réduire les silos entre les directions")
    _attendre_visible(nav, f"#reco-saved-{rid}-objectif", "enregistré")
    _sans_erreur(nav, "Autosave axe et recommandation")

    # Relu en base : valeurs ET drapeau « édité » (la garde de régénération s'y appuie).
    assert _python_sur_base(base_db, _LIRE_RECO, aid, rid) == repr(
        ("Axe gouvernance partagée", True, "Reco initiale",
         "Réduire les silos entre les directions", True))
    # Rechargé : ce que l'utilisateur relit est ce qu'il a tapé.
    nav.naviguer(f"{base}/missions/{mid}/recommandations")
    assert _valeur(nav, f"input[hx-post='/recommandations/axes/{aid}/field']") == "Axe gouvernance partagée"
    assert _valeur(
        nav, f"#reco-card-{rid} textarea[hx-vals*='objectif']") == "Réduire les silos entre les directions"
    assert "Axe gouvernance partagée" in nav.texte()  # synthèse en tête de page
    _sans_erreur(nav, "Recommandations rechargées")


# --------------------------------------------------------------------------- #
# 2. Sauvegarde du brouillon d'enregistrement (audio brut) puis suppression
# --------------------------------------------------------------------------- #
_SEED_ENREGISTREMENT = r"""
from app.db import SessionLocal, init_db
from app.models import Mission, Question, Theme, Trame
init_db()
db = SessionLocal()
m = Mission(name="E2E backup")
db.add(m); db.flush()
tr = Trame(mission_id=m.id); db.add(tr); db.flush()
th = Theme(trame_id=tr.id, title="Gouvernance", position=0); db.add(th); db.flush()
db.add(Question(theme_id=th.id, label="Comment décidez-vous ?", position=0))
db.commit()
print(m.id)
"""


def test_la_sauvegarde_audio_du_brouillon_est_listee_puis_supprimee(
    serveur, nav_micro: Navigateur
) -> None:
    """Faux micro : Démarrer puis Arrêter envoient l'audio brut (POST multipart
    /record/backup). L'écran dit « ✓ sauvegarde audio complète effectuée. », un
    fichier `{mission}_…` existe sur disque, l'onglet Backup de la mission le
    liste (« Non rattachés »), et le clic « Supprimer » (confirm accepté) le
    retire de l'écran ET du disque."""
    nav = nav_micro
    base, dossier = serveur
    base_db = dossier / "e2e.db"
    enregistrements = dossier / "recordings"
    mid = _python_sur_base(base_db, _SEED_ENREGISTREMENT)

    nav.naviguer(f"{base}/missions/{mid}/interviews/record")
    assert not nav.evaluer("document.getElementById('rec-start').disabled"), (
        "Démarrer désactivé : " + nav.evaluer("document.getElementById('rec-status').textContent"))
    nav.cliquer("#rec-start")
    _attendre(lambda: nav.evaluer("!document.getElementById('rec-stop').hidden"),
              "enregistrement jamais démarré (micro refusé ?) : "
              + nav.evaluer("document.getElementById('rec-status').textContent"), 30.0)
    # De l'audio réel à sauvegarder : on attend que le chronomètre avance
    # (état de la page), pas une durée fixe.
    _attendre(lambda: nav.evaluer(
        "(function(){var m=/(\\d+):(\\d+)/.exec(document.getElementById('rec-timer').textContent||'');"
        " return !!m && (parseInt(m[1],10)*60+parseInt(m[2],10))>=3;})()"),
        "le chronomètre d'enregistrement n'avance pas", 30.0)
    nav.cliquer("#rec-stop")
    _attendre_visible(nav, "#rec-backup-status", "sauvegarde audio complète effectuée", 60.0)
    _sans_erreur(nav, "Arrêter l'enregistrement")
    nom = nav.evaluer("document.getElementById('rec-backup-path').value")
    assert nom.startswith(f"{mid}_"), nom
    fichier = enregistrements / nom
    assert fichier.is_file() and fichier.stat().st_size > 1000, (nom, list(enregistrements.glob("*")))
    # Le lecteur de l'écran pointe sur la copie sauvegardée.
    assert nav.evaluer("document.getElementById('rec-audio-download').getAttribute('href')").endswith(
        "/record/backup/" + nom)

    # Onglet Backup de la mission : le fichier y est listé, sous « Non rattachés ».
    nav.naviguer(f"{base}/missions/{mid}")
    nav.cliquer(".tab[data-tab='backup']")
    _attendre(lambda: "Backup (1)" in nav.texte(), "onglet « Backup (1) » jamais affiché : " + nav.texte()[:300])
    panneau = "[data-panel='backup']"
    _attendre_visible(nav, panneau, "Non rattachés")
    _attendre_visible(nav, panneau, nom)
    nav.cliquer(f"{panneau} form[action$='/delete'] button[type=submit]")
    # `cliquer` : le POST + redirection est en vol ; on attend l'état final.
    _attendre(lambda: not fichier.exists(), f"le fichier {nom} est toujours sur le disque")
    nav.naviguer(f"{base}/missions/{mid}")
    nav.cliquer(".tab[data-tab='backup']")
    _sans_erreur(nav, "Supprimer la sauvegarde")
    _attendre(lambda: "Backup (0)" in nav.texte(), "onglet « Backup (0) » jamais affiché")
    _attendre_visible(nav, panneau, "Aucun enregistrement audio pour cette mission.")
    assert nom not in nav.texte()
    assert not list(enregistrements.glob(f"{mid}_*")), list(enregistrements.glob("*"))


# --------------------------------------------------------------------------- #
# 3. Import d'un entretien existant : upload multipart puis confirmation
# --------------------------------------------------------------------------- #
_SEED_IMPORT = r"""
from app.db import SessionLocal, init_db
from app.models import Mission, Question, Theme, Trame
init_db()
db = SessionLocal()
m = Mission(name="E2E import entretien")
db.add(m); db.flush()
tr = Trame(mission_id=m.id); db.add(tr); db.flush()
th = Theme(trame_id=tr.id, title="Gouvernance", position=0); db.add(th); db.flush()
q = Question(theme_id=th.id, label="Comment décidez-vous ?", position=0)
db.add(q); db.commit()
print(m.id, q.id)
"""

_LIRE_IMPORT = r"""
import sys
from app.db import SessionLocal
from app.models import Answer, Interview
db = SessionLocal()
itws = db.query(Interview).filter_by(mission_id=int(sys.argv[1])).all()
print(repr([(i.interviewee_name, i.interviewee_role,
             [(a.text, a.status) for a in i.answers]) for i in itws]))
"""

_REPONSE_FACTICE = "Les décisions se prennent en comité de direction, chaque lundi."


def _docx_entretien(chemin: Path) -> Path:
    from docx import Document

    doc = Document()
    doc.add_paragraph("Compte rendu d'entretien — gouvernance")
    doc.add_paragraph("Les décisions se prennent en comité de direction, chaque lundi matin.")
    doc.save(str(chemin))
    return chemin


def _reponses_saisies(nav: Navigateur) -> str:
    return nav.evaluer(
        "Array.from(document.querySelectorAll('textarea[name=text]'))"
        ".map(function(e){return e.value;}).join(' | ')")


def _envoyer_le_fichier(nav: Navigateur, mid: str, fichier: Path, nom: str, base: str) -> None:
    nav.naviguer(f"{base}/missions/{mid}/interviews/import")
    assert "Importer un entretien depuis un document" in nav.texte()
    form = f"form[action='/missions/{mid}/interviews/import']"
    nav.remplir(f"{form} input[name=interviewee_name]", nom)
    nav.remplir(f"{form} input[name=interviewee_role]", "Directrice des opérations")
    nav.choisir_fichier(f"{form} input[type=file]", fichier)
    nav.cliquer_et_attendre(f"{form} button[type=submit]")


def test_import_entretien_sous_la_coupe_dit_pourquoi_et_ne_cree_rien(serveur, nav: Navigateur,
                                                                   tmp_path) -> None:
    """Chemin réel sous la coupe : l'upload multipart aboutit, l'extraction IA
    échoue, l'écran d'import REVIENT avec le bandeau d'échec lisible ET les
    champs d'identité déjà saisis ; aucun entretien n'est créé."""
    base, dossier = serveur
    base_db = dossier / "e2e.db"
    mid, _qid = _python_sur_base(base_db, _SEED_IMPORT).split()
    docx = _docx_entretien(tmp_path / "entretien.docx")

    _envoyer_le_fichier(nav, mid, docx, "Camille Bernard", base)
    _sans_erreur(nav, "Importer sous la coupe")
    assert "Revue avant import" not in nav.texte()
    _assert_echec_ia_rendu(nav)
    assert _valeur(nav, "input[name=interviewee_name]") == "Camille Bernard"
    assert _valeur(nav, "input[name=interviewee_role]") == "Directrice des opérations"
    assert _python_sur_base(base_db, _LIRE_IMPORT, mid) == "[]"


_LANCEUR_EXTRACTION_FACTICE = r"""
import os, sys
import uvicorn
import app.main
from app.services import interview_extract_ai

def _faux(system, prompt, schema, hint, **kw):
    return {"answers": [{"question_id": int(os.environ["E2E_QID"]),
                         "text": os.environ["E2E_REPONSE"],
                         "verbatims": ["chaque lundi"]}]}

interview_extract_ai.call_ai_json = _faux
uvicorn.run(app.main.app, host="127.0.0.1", port=int(sys.argv[1]), log_level="warning")
"""


@contextlib.contextmanager
def _serveur_extraction_factice(dossier: Path, qid: str):
    """Même uvicorn que `serveur_uvicorn`, sauf que SEULE l'extraction des
    réponses d'un entretien importé (`interview_extract_ai.call_ai_json`) rend
    une réponse canonique : le reste de l'app (routes, templates, base,
    redirections) est le vrai. Sert le chemin heureux que la coupe Ollama
    interdit."""
    port = port_libre()
    env = dict(os.environ, APP_DB_PATH=str(dossier / "e2e.db"), PYTHONUTF8="1",
               WHISPER_WARM_UP="0", OLLAMA_WARM_UP="0",
               OLLAMA_HOST="http://127.0.0.1:9", E2E_QID=qid,
               E2E_REPONSE=_REPONSE_FACTICE)
    journal = open(dossier / "uvicorn-factice.log", "w", encoding="utf-8")
    proc = subprocess.Popen([sys.executable, "-c", _LANCEUR_EXTRACTION_FACTICE, str(port)],
                            cwd=_RACINE, env=env, stdout=journal, stderr=subprocess.STDOUT)
    base = f"http://127.0.0.1:{port}"
    try:
        fin = time.monotonic() + 60
        while time.monotonic() < fin:
            assert proc.poll() is None, (dossier / "uvicorn-factice.log").read_text(
                encoding="utf-8", errors="replace")
            try:
                requete = urllib.request.Request(
                    base + "/missions",
                    headers={"Authorization": "Bearer " + os.environ["APP_AUTH_PASSWORD"]})
                with urllib.request.urlopen(requete, timeout=2):
                    break
            except Exception:
                time.sleep(0.3)
        else:
            raise AssertionError("le serveur factice n'a pas démarré")
        yield base
    finally:
        _arreter(proc)
        journal.close()


def test_import_entretien_revue_puis_validation_creent_l_entretien(
    tmp_path_factory: pytest.TempPathFactory, nav: Navigateur, tmp_path
) -> None:
    """Deux temps : (1) upload du .docx -> écran « Revue avant import » avec la
    question et la réponse proposée ; (2) « Valider l'import » -> fiche de
    l'entretien créé, réponse marquée « à valider ». Une question décochée n'est
    pas importée. Relu en base."""
    dossier = tmp_path_factory.mktemp("e2e-import")
    base_db = dossier / "e2e.db"
    mid, qid = _python_sur_base(base_db, _SEED_IMPORT).split()
    docx = _docx_entretien(tmp_path / "entretien.docx")
    with _serveur_extraction_factice(dossier, qid) as base:
        _envoyer_le_fichier(nav, mid, docx, "Camille Bernard", base)
        _sans_erreur(nav, "Envoyer le .docx")
        texte = nav.texte()
        assert "Revue avant import" in texte
        assert "Comment décidez-vous ?" in texte and _REPONSE_FACTICE in texte, texte[:500]
        # Rien n'est encore écrit : « L'entretien n'est créé qu'à la validation ».
        assert _python_sur_base(base_db, _LIRE_IMPORT, mid) == "[]"

        nav.cliquer_et_attendre("form[action$='/interviews/import/confirm'] button[type=submit]")
        _sans_erreur(nav, "Valider l'import")
        assert "/interviews/" in nav.url() and "/import" not in nav.url(), nav.url()
        texte = nav.texte()
        assert _valeur(nav, "input.identity-name") == "Camille Bernard"
        assert "Comment décidez-vous ?" in texte, texte[:500]
        # La réponse vit dans un champ de saisie (innerText ne la voit pas).
        assert _REPONSE_FACTICE in _reponses_saisies(nav)
        assert _python_sur_base(base_db, _LIRE_IMPORT, mid) == repr(
            [("Camille Bernard", "Directrice des opérations", [(_REPONSE_FACTICE, "to_review")])])
        # Rechargée, la fiche relit la même réponse.
        nav.naviguer(nav.url())
        assert _REPONSE_FACTICE in _reponses_saisies(nav)


def test_import_entretien_question_decochee_n_est_pas_importee(
    tmp_path_factory: pytest.TempPathFactory, nav: Navigateur, tmp_path
) -> None:
    """Le second temps respecte la revue : la case de la question décochée par
    un VRAI clic -> l'entretien est créé SANS cette réponse."""
    dossier = tmp_path_factory.mktemp("e2e-import-decoche")
    base_db = dossier / "e2e.db"
    mid, qid = _python_sur_base(base_db, _SEED_IMPORT).split()
    docx = _docx_entretien(tmp_path / "entretien.docx")
    with _serveur_extraction_factice(dossier, qid) as base:
        _envoyer_le_fichier(nav, mid, docx, "Dana Petit", base)
        assert "Revue avant import" in nav.texte()
        nav.cliquer("input[name=keep]")  # décocher
        assert not nav.evaluer("document.querySelector('input[name=keep]').checked")
        nav.cliquer_et_attendre("form[action$='/interviews/import/confirm'] button[type=submit]")
        _sans_erreur(nav, "Valider sans la question")
        assert _valeur(nav, "input.identity-name") == "Dana Petit"
        assert _REPONSE_FACTICE not in _reponses_saisies(nav)
        assert _python_sur_base(base_db, _LIRE_IMPORT, mid) == repr(
            [("Dana Petit", "Directrice des opérations", [])])


# --------------------------------------------------------------------------- #
# 4. Trames : créer, éditer, supprimer, importer
# --------------------------------------------------------------------------- #
_SEED_TRAME_VIDE = r"""
from app.db import SessionLocal, init_db
from app.models import Mission, Trame
init_db()
db = SessionLocal()
m = Mission(name="E2E trame")
db.add(m); db.flush()
db.add(Trame(mission_id=m.id)); db.commit()
print(m.id)
"""

_LIRE_TRAME = r"""
import sys
from app.db import SessionLocal
from app.models import Mission
m = SessionLocal().get(Mission, int(sys.argv[1]))
print(repr([(t.title, [(q.label, q.qtype, q.help_text) for q in t.questions])
            for t in m.trame.themes]))
"""


def test_trame_creer_editer_et_supprimer(serveur, nav: Navigateur) -> None:
    """Un thème puis une question créés par les formulaires, la question
    retapée au clavier (autosave « ✓ enregistré », type passé en « Échelle »),
    rechargement, puis suppression de la question et du thème (confirm
    accepté). Chaque étape relue en base."""
    base, dossier = serveur
    base_db = dossier / "e2e.db"
    mid = _python_sur_base(base_db, _SEED_TRAME_VIDE)

    nav.naviguer(f"{base}/missions/{mid}/trame")
    _sans_erreur(nav, "Ouvrir la trame")
    nav.remplir("form[action$='/trame/themes'] input[name=title]", "Pilotage de la transformation")
    nav.cliquer_et_attendre("form[action$='/trame/themes'] button[type=submit]")
    assert "1. Pilotage de la transformation" in nav.texte()
    assert "Aucune question dans ce thème." in nav.texte()

    nav.remplir("form.add-question input[name=label]", "Qui arbitre les priorités ?")
    nav.cliquer_et_attendre("form.add-question button[type=submit]")
    assert _valeur(nav, "form.edit-question input[name=label]") == "Qui arbitre les priorités ?"
    assert "Aucune question dans ce thème." not in nav.texte()
    assert _python_sur_base(base_db, _LIRE_TRAME, mid) == repr(
        [("Pilotage de la transformation", [("Qui arbitre les priorités ?", "open", None)])])

    # Édition : intitulé retapé, aide ajoutée, type changé -> autosave.
    _saisir_au_clavier(nav, "form.edit-question input[name=label]", "Qui arbitre les priorités du programme ?")
    _attendre_visible(nav, "form.edit-question .saved", "enregistré")
    nav.evaluer("document.querySelector('form.edit-question .saved').textContent=''")
    _saisir_au_clavier(nav, "form.edit-question input[name=help_text]", "Relancer sur le comité de pilotage")
    _attendre_visible(nav, "form.edit-question .saved", "enregistré")
    _sans_erreur(nav, "Autosave de la question")
    assert _python_sur_base(base_db, _LIRE_TRAME, mid) == repr(
        [("Pilotage de la transformation",
          [("Qui arbitre les priorités du programme ?", "open", "Relancer sur le comité de pilotage")])])
    # Type passé en « Échelle » par la liste (touches réelles) : autosave sur `change`.
    nav.evaluer("document.querySelector('form.edit-question .saved').textContent=''")
    nav.choisir_option("form.edit-question select[name=qtype]", "scale")
    _attendre_visible(nav, "form.edit-question .saved", "enregistré")
    _sans_erreur(nav, "Changer le type de la question")
    _attendre(lambda: _python_sur_base(base_db, _LIRE_TRAME, mid) == repr(
        [("Pilotage de la transformation",
          [("Qui arbitre les priorités du programme ?", "scale", "Relancer sur le comité de pilotage")])]),
        "type « scale » jamais relu en base : " + _python_sur_base(base_db, _LIRE_TRAME, mid))
    nav.naviguer(f"{base}/missions/{mid}/trame")
    assert _valeur(nav, "form.edit-question input[name=label]") == "Qui arbitre les priorités du programme ?"
    assert _valeur(nav, "form.edit-question input[name=help_text]") == "Relancer sur le comité de pilotage"
    assert _valeur(nav, "form.edit-question select[name=qtype]") == "scale"

    # L'aperçu montre ce que l'utilisateur a bâti.
    nav.cliquer_et_attendre(f"a[href='/missions/{mid}/trame/preview']")
    assert "Qui arbitre les priorités du programme ?" in nav.texte()
    assert "1 thème · 1 question" in nav.texte()

    # Suppression : la question, puis le thème.
    nav.naviguer(f"{base}/missions/{mid}/trame")
    nav.cliquer_et_attendre("li.q-edit form.inline[action*='/trame/questions/'][action$='/delete'] button")
    assert "Aucune question dans ce thème." in nav.texte()
    assert "Qui arbitre" not in nav.texte()
    nav.cliquer_et_attendre("form[action*='/trame/themes/'][action$='/delete'] button.btn-danger")
    assert "Pilotage de la transformation" not in nav.texte()
    assert _python_sur_base(base_db, _LIRE_TRAME, mid) == "[]"
    _sans_erreur(nav, "Suppressions")


def _docx_trame(chemin: Path) -> Path:
    from docx import Document

    doc = Document()
    doc.add_paragraph("1. Stratégie et gouvernance")
    doc.add_paragraph("Comment la stratégie est-elle déclinée ?")
    doc.add_paragraph("Qui valide les budgets ?")
    doc.add_paragraph("2. Culture managériale")
    doc.add_paragraph("Comment sont gérés les désaccords ?")
    doc.save(str(chemin))
    return chemin


def test_trame_importer_un_docx_revue_puis_validation(serveur, nav: Navigateur, tmp_path) -> None:
    """Import déterministe (sans IA) : upload du .docx -> « Revue avant import »
    (thèmes et questions détectés), une question décochée, « Valider l'import »
    -> aperçu de la trame avec les seules questions retenues. Relu en base."""
    base, dossier = serveur
    base_db = dossier / "e2e.db"
    mid = _python_sur_base(base_db, _SEED_TRAME_VIDE)
    docx = _docx_trame(tmp_path / "trame.docx")

    nav.naviguer(f"{base}/missions/{mid}/trame")
    nav.cliquer_et_attendre(f"a[href='/missions/{mid}/trame/import']")
    form = f"form[action='/missions/{mid}/trame/import']"
    nav.choisir_fichier(f"{form} input[type=file]", docx)
    nav.cliquer_et_attendre(f"{form} button[type=submit]")
    _sans_erreur(nav, "Envoyer la trame")
    texte = nav.texte()
    assert "Revue avant import" in texte
    for attendu in ("Stratégie et gouvernance", "Culture managériale",
                    "Comment la stratégie est-elle déclinée ?", "Qui valide les budgets ?",
                    "Comment sont gérés les désaccords ?"):
        assert attendu in texte, (attendu, texte[:500])
    assert _python_sur_base(base_db, _LIRE_TRAME, mid) == "[]", "écrit avant la validation"

    nav.cliquer("input[name=keep][value='0-1']")  # « Qui valide les budgets ? »
    nav.cliquer_et_attendre("form[action$='/trame/import/confirm'] button[type=submit]")
    _sans_erreur(nav, "Valider l'import de trame")
    assert nav.url().rstrip("/").endswith(f"/missions/{mid}/trame/preview"), nav.url()
    texte = nav.texte()
    assert "2 thèmes · 2 questions" in texte, texte[:400]
    assert "Comment la stratégie est-elle déclinée ?" in texte
    assert "Qui valide les budgets ?" not in texte
    assert _python_sur_base(base_db, _LIRE_TRAME, mid) == repr(
        [("Stratégie et gouvernance", [("Comment la stratégie est-elle déclinée ?", "open", None)]),
         ("Culture managériale", [("Comment sont gérés les désaccords ?", "open", None)])])
    # Rechargé : l'éditeur relit la trame importée.
    nav.naviguer(f"{base}/missions/{mid}/trame")
    assert "2. Culture managériale" in nav.texte()
