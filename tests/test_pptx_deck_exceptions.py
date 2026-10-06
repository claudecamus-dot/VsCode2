"""Exception contract of the narrowed excepts of app/services/pptx_deck.py.

An EXPECTED error yields the fallback; an UNEXPECTED one (RuntimeError) propagates."""

import pytest

from app.services import pptx_deck as D


class _Ombre:
    def __init__(self, exc):
        self._exc = exc

    @property
    def inherit(self):
        return True

    @inherit.setter
    def inherit(self, v):
        raise self._exc("boom")


class _FormeOmbre:
    def __init__(self, exc):
        self.shadow = _Ombre(exc)


class _FormeGeo:
    def __init__(self, exc):
        self._exc = exc

    @property
    def left(self):
        raise self._exc("boom")


class _Adj:
    def __init__(self, exc):
        self._exc = exc

    def __setitem__(self, i, v):
        raise self._exc("boom")


class _Prs:
    def __init__(self, exc):
        self._exc = exc

    @property
    def slide_masters(self):
        raise self._exc("boom")


@pytest.mark.parametrize("exc", [AttributeError, NotImplementedError])
def test_no_shadow_expected_falls_back(exc):
    D._no_shadow(_FormeOmbre(exc))


def test_no_shadow_unexpected_propagates():
    with pytest.raises(RuntimeError):
        D._no_shadow(_FormeOmbre(RuntimeError))


@pytest.mark.parametrize("exc", [AttributeError, NotImplementedError, ValueError, TypeError])
def test_bornes_in_expected_returns_none(exc):
    assert D._bornes_in(_FormeGeo(exc)) is None


def test_bornes_in_unexpected_propagates():
    with pytest.raises(RuntimeError):
        D._bornes_in(_FormeGeo(RuntimeError))


@pytest.mark.parametrize("exc", [IndexError, AttributeError, KeyError])
def test_police_marque_expected_returns_none(exc):
    assert D.police_marque(_Prs(exc)) is None


def test_police_marque_unexpected_propagates():
    with pytest.raises(RuntimeError):
        D.police_marque(_Prs(RuntimeError))


def test_xfrm_groupe_unexpected_propagates():
    class _El:
        def find(self, *_):
            raise RuntimeError("boom")

    class _Shp:
        _element = _El()

    with pytest.raises(RuntimeError):
        D._xfrm_de_groupe_non_transforme(_Shp())


def test_xfrm_groupe_expected_false():
    class _Shp:
        pass  # no _element -> AttributeError

    assert D._xfrm_de_groupe_non_transforme(_Shp()) is False


# --- Exception classes confirmed by execution on a real python-pptx deck -----
# IndexError   : prs.slide_masters[0] with no master ("slide master index out of range")
# KeyError     : part.part_related_by(RT.THEME) with no theme relation
# IndexError / ValueError : shape.adjustments[i] = v (index out of range / non-numeric)
# NotImplementedError : GraphicFrame.shadow ; TypeError : non-integral xfrm offset

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.util import Inches


def _prs_sans_master():
    prs = Presentation()
    for e in list(prs.part._element.sldMasterIdLst):
        prs.part._element.sldMasterIdLst.remove(e)
    return prs


def test_real_missing_master_raises_indexerror_and_falls_back():
    prs = _prs_sans_master()
    with pytest.raises(IndexError):
        prs.slide_masters[0]
    assert D.police_marque(prs) is None
    D.police_theme(prs)  # falls back to the package scan, no exception escapes
    assert D.theme_colors(prs) == {}


def test_real_missing_theme_relation_raises_keyerror():
    prs = Presentation()
    with pytest.raises(KeyError):
        prs.slides.add_slide(prs.slide_layouts[6]).part.part_related_by(RT.THEME)


def test_police_theme_falls_back_on_missing_relation_via_package_scan():
    class _Master:
        class part:
            @staticmethod
            def part_related_by(_):
                raise KeyError("no theme")

    class _Part:
        partname = "/ppt/theme/theme1.xml"
        blob = b'<a:minorFont><a:latin typeface="Arial"/></a:minorFont>'

    class _Pkg:
        @staticmethod
        def iter_parts():
            return [_Part]

    class _P:
        slide_masters = [_Master]

        class part:
            package = _Pkg

    assert D.police_theme(_P) == "Arial"


@pytest.mark.parametrize("exc", [AttributeError, TypeError])
def test_police_theme_scan_expected_returns_none(exc):
    class _Master:
        class part:
            @staticmethod
            def part_related_by(_):
                raise KeyError("no theme")

    class _Pkg:
        @staticmethod
        def iter_parts():
            raise exc("boom")

    class _P:
        slide_masters = [_Master]

        class part:
            package = _Pkg

    assert D.police_theme(_P) is None


def test_police_theme_scan_unexpected_propagates():
    class _Master:
        class part:
            @staticmethod
            def part_related_by(_):
                raise KeyError("no theme")

    class _Pkg:
        @staticmethod
        def iter_parts():
            raise RuntimeError("boom")

    class _P:
        slide_masters = [_Master]

        class part:
            package = _Pkg

    with pytest.raises(RuntimeError):
        D.police_theme(_P)


def test_police_theme_first_lookup_unexpected_propagates():
    class _Master:
        class part:
            @staticmethod
            def part_related_by(_):
                raise RuntimeError("boom")

    class _P:
        slide_masters = [_Master]

    with pytest.raises(RuntimeError):
        D.police_theme(_P)


def test_add_forme_bad_adjustments_fall_back_real_shape():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    prst = next(iter(D.FORMES_PRST))
    # out-of-range index and non-numeric value: both really raise, neither escapes
    D.add_forme(slide, prst, 0, 0, 1, 1, adj=[0.1] * 9)
    D.add_forme(slide, prst, 0, 0, 1, 1, adj=["x"])


def test_add_forme_adjustments_unexpected_propagates(monkeypatch):
    class _Shp:
        adjustments = _Adj(RuntimeError)
        rotation = 0

        def __getattr__(self, name):
            return lambda *a, **k: None

    class _Shapes:
        @staticmethod
        def add_shape(*a, **k):
            return _Shp()

    class _Slide:
        shapes = _Shapes

    monkeypatch.setattr(D, "_no_shadow", lambda s: None)
    prst = next(iter(D.FORMES_PRST))
    with pytest.raises(RuntimeError):
        D.add_forme(_Slide, prst, 0, 0, 1, 1, adj=[0.1])


def test_theme_colors_external_target_valueerror_falls_back():
    class _Rel:
        reltype = RT.THEME

        @property
        def target_part(self):
            raise ValueError("external")

    class _P:
        class _M:
            class part:
                rels = {"r": _Rel()}

        slide_masters = [_M]

    assert D.theme_colors(_P) == {}


def test_theme_colors_unexpected_propagates():
    class _Rel:
        reltype = RT.THEME

        @property
        def target_part(self):
            raise RuntimeError("boom")

    class _P:
        class _M:
            class part:
                rels = {"r": _Rel()}

        slide_masters = [_M]

    with pytest.raises(RuntimeError):
        D.theme_colors(_P)


def test_theme_colors_without_colour_node_returns_empty_dict():
    class _TP:
        blob = b"<a:theme></a:theme>"

    class _Rel:
        reltype = RT.THEME
        target_part = _TP

    class _P:
        class _M:
            class part:
                rels = {"r": _Rel()}

        slide_masters = [_M]

    assert D.theme_colors(_P) == {}


class _ShpGeo:
    name = "s"

    def __init__(self, exc):
        self._exc = exc

    @property
    def left(self):
        raise self._exc("boom")


class _PrsGeo:
    slide_width = 100
    slide_height = 100

    def __init__(self, exc):
        class _S:
            shapes = [_ShpGeo(exc)]
        self.slides = [_S()]


@pytest.mark.parametrize("exc", [AttributeError, NotImplementedError, ValueError, TypeError])
def test_verifier_geometrie_expected_skips_shape(exc):
    assert D.verifier_geometrie(_PrsGeo(exc)) == []


def test_verifier_geometrie_unexpected_propagates():
    with pytest.raises(RuntimeError):
        D.verifier_geometrie(_PrsGeo(RuntimeError))


def test_verifier_geometrie_real_graphicframe_and_bad_xfrm_do_not_escape():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_table(1, 1, 0, 0, Inches(1), Inches(1))
    D.verifier_geometrie(prs)
