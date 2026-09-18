"""Filet de geometrie de pptx_deck.py (finding hub:pptx-deck/verifier_geometrie,
salle conseil-flotte du 2026-09-18, portage depuis VSCode3/VScode5) :
`verifier_geometrie` ne testait que les bords des formes. Une boite inversee
(hauteur ou largeur negative) a son bord bas au-dessus de son bord haut : elle
ne depasse aucun bord, le filet rendait `[]` et le deck s'ecrivait sous
"CONTROLE: OK" alors que PowerPoint refuse d'ouvrir un fichier portant une
telle forme.

Chaque test est ADVERSARIAL : il prouve d'abord que la boite posee est
reellement inversee/nulle avant de faire confiance au verdict du filet.
"""
from __future__ import annotations

from pptx import Presentation
from pptx.util import Inches

from app.services.pptx_deck import add_rect, add_text, verifier_geometrie


def _prs_vide() -> Presentation:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    prs.slides.add_slide(prs.slide_layouts[6])
    return prs


def test_hauteur_negative_est_signalee() -> None:
    prs = _prs_vide()
    shp = add_rect(prs.slides[0], 1.0, 1.0, 3.0, -0.4, fill="#ffffff")
    assert shp.height < 0, "le cas de test doit vraiment poser une boite inversee"
    l, t, w, h = shp.left, shp.top, shp.width, shp.height
    assert 0 <= l and 0 <= t and (l + w) <= prs.slide_width \
        and (t + h) <= prs.slide_height, \
        "cette boite ne depasse aucun bord : seul un controle de dimension la voit"

    problemes = verifier_geometrie(prs)

    assert len(problemes) == 1, problemes
    assert "dimension non positive" in problemes[0]
    assert "h=-0.40" in problemes[0]


def test_hauteur_nulle_est_signalee() -> None:
    prs = _prs_vide()
    add_rect(prs.slides[0], 1.0, 1.0, 3.0, 0.0, fill="#ffffff")

    problemes = verifier_geometrie(prs)

    assert len(problemes) == 1, problemes
    assert "dimension non positive" in problemes[0]


def test_largeur_negative_est_signalee() -> None:
    prs = _prs_vide()
    add_rect(prs.slides[0], 4.0, 1.0, -2.0, 0.5, fill="#ffffff")

    problemes = verifier_geometrie(prs)

    assert len(problemes) == 1, problemes
    assert "dimension non positive" in problemes[0]
    assert "w=-2.00" in problemes[0]


def test_hauteur_calculee_par_soustraction_qui_decroche_est_signalee() -> None:
    """Scenario realiste : une hauteur calculee par soustraction (bas fixe -
    haut variable) quand le haut a deborde sous le bas."""
    prs = _prs_vide()
    bas_fixe = 5.45
    top = bas_fixe + 0.3          # le contenu a deborde sous le plancher
    card_h = bas_fixe - top       # -> -0.30in
    add_rect(prs.slides[0], 0.615, top, 3.0, card_h, fill="#ffffff")

    assert [p for p in verifier_geometrie(prs) if "dimension non positive" in p]


def test_formes_correctes_ne_declenchent_rien() -> None:
    """Le pendant obligatoire : un filet qui crie sur une slide correcte est
    debranche au premier build. Trois formes nominales, zero constat."""
    prs = _prs_vide()
    add_rect(prs.slides[0], 0.615, 0.5, 3.0, 1.2, fill="#ffffff")
    add_rect(prs.slides[0], 0.615, 2.0, 8.0, 0.02, fill="#ffffff")  # filet fin
    add_text(prs.slides[0], 0.615, 3.0, 3.0, 0.6, [("court", {})])

    assert verifier_geometrie(prs) == []
