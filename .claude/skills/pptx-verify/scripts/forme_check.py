"""Contrôle de FORME d'un .pptx, déterministe (0 token) : avertit, ne bloque jamais par défaut.

Ce que le contrôle voit : formes hors slide, image étirée, image sans texte alternatif, mention de travail
(« à valider », « V0 »…), police explicite trop petite, mur de texte, titre qui n'est pas un placeholder de titre, trop de tailles
de police sur une slide, slide vide, notes d'orateur absentes (information). Le texte des tableaux compte ; la
géométrie des enfants d'un groupe n'est pas contrôlée (repère interne du groupe). Ce qu'il ne voit PAS :
contraste, désalignement, composition, message. Il ne remplace donc pas le rendu regardé
(pptx-verify, étape 2) : un deck sans avertissement peut rester illisible.

Les seuils sont dans forme_check_seuils.json à côté du script ; tant que `etalonne` y vaut false,
ils sont PROVISOIRES (valeurs de départ, pas un jugement de qualité) et le rapport le dit.
Les tailles de police héritées du gabarit ne sont pas lues : le contrôle de taille est un plancher.

Usage :
  py forme_check.py deck.pptx [--json] [--strict] [--seuils fichier.json]
  py forme_check.py --distribution deck1.pptx dossier/ ...   (mesures par slide, pour étalonner)
Code de retour : 0 (même avec avertissements), 1 avec --strict et au moins un avertissement,
2 si python-pptx est absent ou le fichier illisible.
"""
import argparse
import json
import os
import re
import sys

SEUILS_DEFAUT = {"police_min_pt": 10, "caracteres_max_slide": 900, "tailles_max_slide": 6,
                 "ratio_image_tolerance": 0.05, "etalonne": False}
# Densité, petite police et nombre de tailles sont en INFORMATION : sur 3 decks jugés par l'utilisateur
# (2 bons, 1 mauvais), ces mesures vont à l'envers (les bons decks sont plus denses, à police plus petite)
# et ne départagent rien ; voir docs/redaction/etalonnage-forme-decks.md. Ils restent affichés.
SEVERITE = {"SANS_NOTES": "info", "POLICE_PETITE": "info", "MUR_DE_TEXTE": "info", "TROP_DE_TAILLES": "info"}
MARQUEUR_DE_TRAVAIL = re.compile(r"\b(?:à|a) (?:valider|calibrer|confirmer)\b|\bTODO\b|\blorem ipsum\b|"
                                 r"\bversion de travail\b|\bV0\b", re.I)
TITRES = ("TITLE", "CENTER_TITLE")
# Texte de gabarit laissé dans un livrable : vu sur 2 decks sur 3 (revue du 2026-10-13 : Pitch, 12 slides sur 16
# signalées ; Offre, 11 sur 32), invisible pour les mesures de densité (les formes du gabarit existent).
# Motifs resserrés après le comité du même jour : « WIP » et « à ajouter » ne comptent qu'en marque de fin
# (« Offre - WIP », « Définition produit à ajouter »), pas dans « Limiter le WIP » ni « Contacts à ajouter au CRM ».
GABARIT_RESIDUEL = re.compile(
    r"JJ\.MM\.\d\d|NOM DU CLIENT|\(calcul[ée]\)|\(voir menu\)|\bà ajouter\s*[.:]?\s*$|"
    r"[-–—(\[]\s*WIP\s*[)\]]?\s*$|^\s*WIP\s*$|"
    r"ici mettre une photo|template de photo|remplacer l.image|cliquez pour modifier|click to edit|"
    r"\bvotre titre ici\b|\bTBD\b|\bX{3,}\b|\bà (?:définir|compléter|rédiger)\b|"
    r"étapes,\s*activités clés,\s*livrables|\bnote interne\b|\bà supprimer\b|\bbrouillon\b", re.I)
# Paragraphe qui est À LUI SEUL un libellé de gabarit (« Texte », « Titre », « ... ») : un mot isolé n'est pas
# un motif de phrase, il ne vaut résidu que s'il remplit tout le paragraphe.
GABARIT_PARAGRAPHE = re.compile(r"^\s*(?:texte|titre|sous-titre|lorem ipsum.*|\.{3}|…)\s*$", re.I)
PLACEHOLDERS_TECHNIQUES = ("DATE", "FOOTER", "SLIDE_NUMBER")
# « (suite) » ou « - suite » en fin de titre ; pas « Et maintenant, la suite » (une suite de phrase, pas un sujet coupé).
SUITE_EN_TITRE = re.compile(r"\(suite(?:\s*\d+)?\)|[-–—:]\s*suite(?:\s*\d+)?\s*$", re.I)
# Layouts d'intercalaire, de couverture ou de clôture : un titre seul y est voulu, pas un corps oublié.
LAYOUT_TITRE_SEUL = re.compile(r"section|intercalaire|chapitre|couverture|title slide|titre de section|merci|clôture|cloture|fin\b|"
                               r"title only|titre seul", re.I)
# Titres qui reviennent légitimement d'une slide à l'autre.
TITRES_REPETABLES = re.compile(r"^(?:annexes?|merci|questions?|sommaire|agenda)\b", re.I)
NOM_DE_FICHIER = re.compile(r"\.(png|jpe?g|gif|bmp|svg|emf|wmf|tiff?)$", re.I)


def charger_seuils(chemin=None):
    """Seuils par défaut, surchargés par le JSON voisin (ou `chemin`) ; un JSON illisible est ignoré."""
    seuils = dict(SEUILS_DEFAUT)
    chemin = chemin or os.path.join(os.path.dirname(os.path.abspath(__file__)), "forme_check_seuils.json")
    try:
        with open(chemin, encoding="utf-8") as fh:
            lu = json.load(fh)
        if isinstance(lu, dict):
            seuils.update({k: v for k, v in lu.items() if k in SEUILS_DEFAUT})
    except (OSError, ValueError):
        pass
    return seuils


def _formes(shapes):
    """Toutes les formes, groupes inclus."""
    for sh in shapes:
        yield sh
        try:
            if hasattr(sh, "shapes"):
                yield from _formes(sh.shapes)
        except Exception:  # noqa: BLE001 - un groupe illisible ne fait pas tomber le contrôle
            continue


def _est_titre(sh):
    try:
        # nom exact de l'enum : « SUBTITLE » contient « TITLE », une sous-chaîne le prendrait pour un titre
        return bool(sh.is_placeholder) and getattr(sh.placeholder_format.type, "name", "") in TITRES
    except Exception:  # noqa: BLE001
        return False


def _paragraphes(sh):
    """Paragraphes d'une forme : zone de texte ou cellules de tableau ; vide si illisible."""
    try:
        if getattr(sh, "has_text_frame", False) and sh.text_frame.text.strip():
            return list(sh.text_frame.paragraphs)
        if getattr(sh, "has_table", False) and sh.has_table:
            return [p for row in sh.table.rows for cell in row.cells for p in cell.text_frame.paragraphs]
    except Exception:  # noqa: BLE001
        pass
    return []


def _est_contenu_visuel(sh):
    """Image, tableau, graphique ou schéma natif : une slide qui en porte un n'est pas un corps vide."""
    try:  # shape_type couvre déjà tableau et graphique : les branches has_table/has_chart étaient redondantes
        return sh.shape_type is not None and any(k in str(sh.shape_type) for k in ("PICTURE", "CHART", "TABLE", "MEDIA"))
    except Exception:  # noqa: BLE001
        return False


def _placeholder_texte_vide(sh):
    """Placeholder de texte (hors date/pied/numéro) resté vide : le gabarit n'a pas été rempli."""
    try:
        return (bool(sh.is_placeholder) and getattr(sh, "has_text_frame", False)
                and getattr(sh.placeholder_format.type, "name", "") not in PLACEHOLDERS_TECHNIQUES
                and not sh.text_frame.text.strip())
    except Exception:  # noqa: BLE001
        return False


def _layout_titre_seul(slide):
    """Vrai si le nom du layout est celui d'un intercalaire, d'une couverture ou d'une clôture."""
    try:
        return bool(LAYOUT_TITRE_SEUL.search(slide.slide_layout.name or ""))
    except Exception:  # noqa: BLE001
        return False


def mesures_slide(slide, largeur, hauteur):
    """Mesures brutes d'une slide : caractères, plus petite taille explicite, tailles distinctes, titre."""
    chars, tailles, titre, n, marqueur = 0, set(), False, 0, False
    residuel, vides, visuel, titre_texte = [], 0, False, ""
    corps, dessin = 0, 0
    for sh in _formes(slide.shapes):
        n += 1
        est_titre = _est_titre(sh)
        titre = titre or est_titre
        if est_titre and not titre_texte:
            titre_texte = " ".join(sh.text_frame.text.split())
        vides += _placeholder_texte_vide(sh)
        visuel = visuel or _est_contenu_visuel(sh)
        # Illustration dessinée en formes natives : 3 formes non-placeholder sans texte ou plus (une seule est
        # un filet ou un cadre de gabarit, pas un schéma).
        try:
            if not sh.is_placeholder and not getattr(sh, "has_table", False) and "GROUP" not in str(sh.shape_type) \
                    and not (getattr(sh, "has_text_frame", False) and sh.text_frame.text.strip()):
                dessin += 1
        except Exception:  # noqa: BLE001
            pass
        for para in _paragraphes(sh):
            chars += len(para.text)
            if not est_titre:
                corps += len(para.text.strip())
            marqueur = marqueur or bool(MARQUEUR_DE_TRAVAIL.search(para.text))
            trouve = GABARIT_RESIDUEL.search(para.text) or GABARIT_PARAGRAPHE.search(para.text)
            if trouve and trouve.group(0).strip() not in residuel:
                residuel.append(trouve.group(0).strip())
            for run in para.runs:
                try:
                    if run.font.size:
                        tailles.add(run.font.size.pt)
                except Exception:  # noqa: BLE001
                    continue
    return {"caracteres": chars, "taille_min_pt": min(tailles) if tailles else None,
            "tailles_distinctes": len(tailles), "titre_placeholder": titre, "formes": n, "marqueur_de_travail": marqueur,
            "gabarit_residuel": residuel, "placeholders_vides": vides,
            "contenu_visuel": visuel or dessin >= 3, "titre_texte": titre_texte, "caracteres_corps": corps,
            "layout_titre_seul": _layout_titre_seul(slide)}


def analyser(prs, seuils=None):
    """Liste de constats {slide, code, severite, forme, detail} ; ne lève pas sur une forme illisible."""
    s = seuils or charger_seuils()
    W, H, out = prs.slide_width, prs.slide_height, []

    def ajouter(i, code, forme, detail):
        out.append({"slide": i, "code": code, "severite": SEVERITE.get(code, "avertissement"),
                    "forme": forme, "detail": detail})

    titres_vus = {}
    for i, slide in enumerate(prs.slides, 1):
        m = mesures_slide(slide, W, H)
        if m["formes"] == 0:
            ajouter(i, "SLIDE_VIDE", "-", "aucune forme")
        elif m["caracteres_corps"] == 0 and not m["contenu_visuel"] and not m["layout_titre_seul"]:
            # Rien d'autre que le titre : le texte du titre ne compte pas (supprimer le placeholder de corps ne
            # doit pas faire passer la slide). Un placeholder resté vide est la cause la plus fréquente.
            if m["placeholders_vides"]:
                ajouter(i, "PLACEHOLDER_VIDE", "-", f"{m['placeholders_vides']} placeholder(s) de texte vide(s) et "
                        "aucun autre texte, image, tableau ni schéma sur la slide")
            else:
                ajouter(i, "CORPS_VIDE", "-", "aucun texte hors titre, ni image, tableau, graphique ou schéma")
        if m["gabarit_residuel"]:
            ajouter(i, "GABARIT_RESIDUEL", "-", "texte de gabarit laissé : " + " | ".join(m["gabarit_residuel"]))
        if m["titre_texte"]:
            cle = m["titre_texte"].casefold()
            if cle in titres_vus:
                precedente, intercalaire = titres_vus[cle]
                # Un intercalaire suivi de sa première slide de contenu porte le même titre : pas un doublon.
                if not intercalaire and not TITRES_REPETABLES.search(m["titre_texte"]):
                    ajouter(i, "TITRE_DOUBLON", "-", f"même titre que la slide {precedente} : « {m['titre_texte'][:60]} »")
            else:
                titres_vus[cle] = (i, m["layout_titre_seul"] or m["caracteres_corps"] <= 10)
            if SUITE_EN_TITRE.search(m["titre_texte"]):
                ajouter(i, "SUITE_EN_TITRE", "-", "titre en « (suite) » : un sujet coupé au lieu d'un message par slide")
        if not (slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip()):
            ajouter(i, "SANS_NOTES", "-", "pas de notes d'orateur")
        if m["formes"] and not m["titre_placeholder"]:
            ajouter(i, "TITRE_NON_PLACEHOLDER", "-", "aucun placeholder de titre : plan non navigable")
        if m["marqueur_de_travail"]:
            ajouter(i, "MARQUEUR_DE_TRAVAIL", "-", "mention de travail (« à valider », « V0 », TODO…) dans un livrable")
        if m["caracteres"] > s["caracteres_max_slide"]:
            ajouter(i, "MUR_DE_TEXTE", "-", f"{m['caracteres']} caractères > {s['caracteres_max_slide']}")
        if m["tailles_distinctes"] > s["tailles_max_slide"]:
            ajouter(i, "TROP_DE_TAILLES", "-", f"{m['tailles_distinctes']} tailles > {s['tailles_max_slide']}")
        if m["taille_min_pt"] is not None and m["taille_min_pt"] < s["police_min_pt"]:
            ajouter(i, "POLICE_PETITE", "-", f"{m['taille_min_pt']} pt < {s['police_min_pt']} pt (taille explicite)")
        for sh in slide.shapes:  # niveau slide seulement : les coordonnées d'un enfant de groupe sont dans le repère du groupe
            try:
                if None not in (sh.left, sh.top, sh.width, sh.height) and (
                        sh.left < 0 or sh.top < 0 or sh.left + sh.width > W + 9525 or sh.top + sh.height > H + 9525):
                    ajouter(i, "HORS_SLIDE", sh.name, "dépasse le cadre de la slide")
            except Exception:  # noqa: BLE001
                continue
        for sh in _formes(slide.shapes):
            try:
                if sh.shape_type is not None and "PICTURE" in str(sh.shape_type):
                    iw, ih = sh.image.size
                    ratio = (sh.width / sh.height) / (iw / ih)
                    rogne = (sh.crop_left + sh.crop_right, sh.crop_top + sh.crop_bottom)
                    if rogne == (0, 0) and abs(ratio - 1) > s["ratio_image_tolerance"]:
                        ajouter(i, "IMAGE_DEFORMEE", sh.name, f"ratio forme/image {ratio:.2f} sans recadrage")
                    alt = (sh._element.nvPicPr.cNvPr.get("descr") or "").strip()
                    if not alt or NOM_DE_FICHIER.search(alt):  # python-pptx met le nom du fichier par défaut
                        ajouter(i, "IMAGE_SANS_ALT", sh.name, "pas de texte alternatif" if not alt else "texte alternatif = nom de fichier")
            except Exception:  # noqa: BLE001 - une forme illisible est ignorée, pas fatale
                continue
    return out


def _charger_pptx():
    try:
        from pptx import Presentation  # noqa: PLC0415
    except ImportError:
        print("forme_check : python-pptx est absent (pip install python-pptx).", file=sys.stderr)
        sys.exit(2)
    return Presentation


def _percentile(valeurs, p):
    v = sorted(valeurs)
    return v[min(len(v) - 1, int(round(p / 100 * (len(v) - 1))))] if v else None


def distribution(chemins):
    """Percentiles par slide sur un corpus : de quoi étalonner les seuils, pas un jugement de qualité."""
    Presentation = _charger_pptx()
    fichiers = []
    for c in chemins:
        if os.path.isdir(c):
            for d, _, fs in os.walk(c):
                fichiers += [os.path.join(d, f) for f in fs if f.lower().endswith(".pptx") and not f.startswith("~$")]
        else:
            fichiers.append(c)
    chars, mins, tailles, lus = [], [], [], 0
    for f in fichiers:
        try:
            prs = Presentation(f)
        except Exception:  # noqa: BLE001
            continue
        lus += 1
        for sl in prs.slides:
            m = mesures_slide(sl, prs.slide_width, prs.slide_height)
            chars.append(m["caracteres"])
            tailles.append(m["tailles_distinctes"])
            if m["taille_min_pt"] is not None:
                mins.append(m["taille_min_pt"])
    return {"fichiers_lus": lus, "slides": len(chars),
            "caracteres": {f"p{p}": _percentile(chars, p) for p in (50, 75, 90, 95)},
            "taille_min_pt": {f"p{p}": _percentile(mins, p) for p in (5, 10, 25, 50)},
            "tailles_distinctes": {f"p{p}": _percentile(tailles, p) for p in (50, 75, 90, 95)}}


def main(argv=None):
    for flux in (sys.stdout, sys.stderr):  # accents lisibles quand la sortie est redirigée (cp1252 sous Windows)
        try:
            flux.reconfigure(encoding="utf-8")
        except (AttributeError, ValueError):
            pass
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("chemins", nargs="+")
    ap.add_argument("--json", action="store_true")
    ap.add_argument("--strict", action="store_true", help="code 1 s'il y a au moins un avertissement")
    ap.add_argument("--seuils", default=None)
    ap.add_argument("--distribution", action="store_true")
    a = ap.parse_args(argv)
    if a.distribution:
        print(json.dumps(distribution(a.chemins), ensure_ascii=False, indent=2))
        return 0
    Presentation = _charger_pptx()
    if a.seuils and not os.path.isfile(a.seuils):
        print(f"forme_check : fichier de seuils introuvable ({a.seuils}).", file=sys.stderr)
        return 2
    seuils = charger_seuils(a.seuils)
    resultats, il_y_a_avertissement = [], False
    for chemin in a.chemins:
        try:
            prs = Presentation(chemin)
        except Exception as exc:  # noqa: BLE001
            print(f"forme_check : fichier illisible ({chemin} : {exc}).", file=sys.stderr)
            return 2
        constats = analyser(prs, seuils)
        avert = [c for c in constats if c["severite"] == "avertissement"]
        il_y_a_avertissement = il_y_a_avertissement or bool(avert)
        resultats.append((chemin, constats, avert))
    if a.json:
        sortie = {"seuils": seuils, "constats": resultats[0][1]} if len(resultats) == 1 else {
            "seuils": seuils, "fichiers": [{"fichier": c, "constats": k} for c, k, _ in resultats]}
        print(json.dumps(sortie, ensure_ascii=False, indent=2))
    else:
        if not seuils["etalonne"]:
            print("SEUILS PROVISOIRES (non étalonnés) : un avertissement est un signal, pas un verdict.")
        for chemin, constats, avert in resultats:
            if len(resultats) > 1:
                print(f"== {chemin}")
            for c in constats:
                print(f"slide {c['slide']} | {c['code']} | {c['forme']} | {c['detail']}")
            print(f"{len(avert)} avertissement(s), {len(constats) - len(avert)} information(s).")
        print("Ce contrôle ne voit ni contraste, ni alignement, ni composition : le rendu regardé reste obligatoire.")
    return 1 if (a.strict and il_y_a_avertissement) else 0


if __name__ == "__main__":
    sys.exit(main())
