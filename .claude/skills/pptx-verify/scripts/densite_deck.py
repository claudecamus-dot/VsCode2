"""Per-slide density of a .pptx: characters, text zones, share of text under 10 pt.

NON-REGRESSION CEILING, NOT A CHARTER NORM. The ACN charter gives minimum font sizes but no
character cap per slide type. `regressions(avant, apres)` only says "this slide became denser
than the previous version", never "this slide is too dense". Taken over from Vscode8
(scripts/densite.py); thresholds moved to densite_deck_seuils.json next to this script.

Limits: explicit run font sizes only (inherited sizes are ignored); `regressions` compares
BY SLIDE RANK, so it is only valid for two versions of a deck with the SAME structure (same
slide count and order). Decks with a different slide count are refused (ValueError).

Usage: py densite_deck.py deck.pptx [--top=10]
"""
import json
import os
import sys

SEUILS_DEFAUT = {"ratio": 1.10, "marge_car": 20, "marge_pct_petit": 10, "petit_pt": 10}


def charger_seuils(chemin=None):
    """Defaults overridden by the neighbouring JSON (or `chemin`); an unreadable JSON is ignored."""
    seuils = dict(SEUILS_DEFAUT)
    chemin = chemin or os.path.join(os.path.dirname(os.path.abspath(__file__)), "densite_deck_seuils.json")
    try:
        with open(chemin, encoding="utf-8") as fh:
            lu = json.load(fh)
        if isinstance(lu, dict):
            seuils.update({k: v for k, v in lu.items() if k in SEUILS_DEFAUT})
    except (OSError, ValueError):
        pass
    return seuils


def _taille_pt(r):
    """Explicit run size in pt without creating an rPr (python-pptx's `font.size` would mutate)."""
    rpr = r._r.rPr
    sz = rpr.get("sz") if rpr is not None else None
    return int(sz) / 100 if sz else None


def mesurer(prs, seuils=None):
    """{rank: {car, zones, pct_petit}}; pct_petit = % of characters with explicit size < petit_pt."""
    s = seuils or charger_seuils()
    out = {}
    for i, sl in enumerate(prs.slides, 1):
        car = petit = zones = 0
        for sh in sl.shapes:
            if getattr(sh._element, "txBody", None) is None or not sh.text_frame.text.strip():
                continue
            zones += 1
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    n = len(r.text.strip())
                    car += n
                    if _taille_pt(r) is not None and _taille_pt(r) < s["petit_pt"]:
                        petit += n
        out[i] = {"car": car, "zones": zones, "pct_petit": round(100 * petit / car) if car else 0}
    return out


def regressions(avant, apres, seuils=None):
    """Ranks whose characters or small-text share exceed `avant` beyond the thresholds.

    Compares by rank: both measures must have the same slides, else ValueError.
    """
    if set(avant) != set(apres):
        raise ValueError("densite_deck: rank comparison needs decks of the same structure "
                         f"({len(avant)} slides before, {len(apres)} after)")
    s = seuils or charger_seuils()
    return [i for i, m in apres.items()
            if m["car"] > avant[i]["car"] * s["ratio"] + s["marge_car"]
            or m["pct_petit"] > avant[i]["pct_petit"] + s["marge_pct_petit"]]


def main(argv=None):
    from pptx import Presentation
    argv = sys.argv[1:] if argv is None else argv
    if not argv:
        print("usage: densite_deck.py deck.pptx [--top=N]", file=sys.stderr)
        return 2
    top = int(next((a.split("=")[1] for a in argv if a.startswith("--top=")), 10))
    m = mesurer(Presentation(argv[0]))
    for i, v in sorted(m.items(), key=lambda kv: -kv[1]["car"])[:top]:
        print(f"slide {i:3d} : {v['car']:5d} car., {v['zones']:3d} zones, {v['pct_petit']:3d} % sous 10 pt")
    return 0


if __name__ == "__main__":
    sys.exit(main())
