"""Rend en PNG seulement quelques slides d'un deck (les autres sont écartées sur une copie temporaire).

Usage : py rendu_slide.py deck.pptx 7 13 [--dossier=D] [--dpi=80]
Sans --dossier : sous-dossier neuf C:/tmp/rendu/rendu_XXXX. Les PNG portent le rang d'origine, slides masquées comprises (s007.png) ; une slide masquée demandée est
démasquée dans la copie, sinon LibreOffice l'omettrait. Moteur LibreOffice headless avec profil dédié
(évite le verrou d'une instance ouverte). Jamais PowerPoint COM. Police et moteur = ceux de CETTE machine :
un rendu complet reste nécessaire avant livraison. Ne modifie jamais le deck source.
"""
import os
import shutil
import subprocess
import sys
import tempfile

SOFFICE = shutil.which("soffice") or r"C:\Program Files\LibreOffice\program\soffice.exe"


def dossier_frais():
    """Sous-dossier neuf par exécution (le dossier partagé C:/tmp/rendu garde des PNG périmés)."""
    return tempfile.mkdtemp(prefix="rendu_", dir=_racine_rendu())


def _racine_rendu():
    racine = "C:/tmp/rendu" if os.path.isdir("C:/tmp") else tempfile.gettempdir()
    os.makedirs(racine, exist_ok=True)
    return racine


def mini_deck(deck, rangs, sortie):
    """Écrit dans `sortie` une copie de `deck` ne gardant que les slides de rang `rangs` (1-based, masquées incluses)."""
    from pptx import Presentation
    prs = Presentation(deck)
    n = len(prs.slides)
    hors = [r for r in rangs if not 1 <= r <= n]
    if hors:
        raise ValueError(f"rang(s) hors deck ({n} slides) : {hors}")
    ids = prs.slides._sldIdLst
    for i, sld_id in reversed(list(enumerate(list(ids), 1))):
        if i not in rangs:
            prs.part.drop_rel(sld_id.rId)
            ids.remove(sld_id)
    for sl in prs.slides:
        if sl._element.get("show") == "0":
            del sl._element.attrib["show"]
    prs.save(sortie)


def lancer(cmd, timeout):
    """Lance `cmd` ; au dépassement de `timeout`, tue l'ARBRE de NOTRE enfant (taskkill /F /T sur son PID,
    jamais un autre process : soffice.exe n'est que le lanceur de soffice.bin) puis lève TimeoutError."""
    proc = subprocess.Popen(cmd, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    try:
        rc = proc.wait(timeout=timeout)
    except subprocess.TimeoutExpired:
        if os.name == "nt":
            subprocess.run(["taskkill", "/F", "/T", "/PID", str(proc.pid)], capture_output=True, check=False)
        else:
            proc.kill()
        proc.wait()
        raise TimeoutError(f"rendu : {timeout} s dépassées, processus {proc.pid} et enfants arrêtés") from None
    if rc:
        raise subprocess.CalledProcessError(rc, cmd)


def rendre(deck, rangs, dossier, dpi=80):
    import fitz
    rangs = sorted(set(rangs))
    os.makedirs(dossier, exist_ok=True)
    for rang in rangs:  # purge des seuls PNG demandés : aucun reliquat d'un rendu précédent ne passe pour frais
        try:
            os.remove(os.path.join(dossier, f"s{rang:03d}.png"))
        except FileNotFoundError:
            pass
    with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as tmp:
        mini = os.path.join(tmp, "mini.pptx")
        mini_deck(deck, rangs, mini)
        profil = "-env:UserInstallation=file:///" + os.path.join(tmp, "lo").replace("\\", "/")
        lancer([SOFFICE, profil, "--headless", "--convert-to", "pdf", "--outdir", tmp, mini], 180)
        pdf = fitz.open(os.path.join(tmp, "mini.pdf"))
        try:
            if len(pdf) != len(rangs):
                raise RuntimeError(f"rendu : {len(pdf)} page(s) pour {len(rangs)} slide(s) demandée(s)")
            sorties = []
            for page, rang in zip(pdf, rangs, strict=True):
                chemin = os.path.join(dossier, f"s{rang:03d}.png")
                page.get_pixmap(dpi=dpi).save(chemin)
                sorties.append(chemin)
        finally:
            pdf.close()
        return sorties


if __name__ == "__main__":
    opt = dict(a.lstrip("-").split("=", 1) for a in sys.argv[1:] if a.startswith("--") and "=" in a)
    pos = [a for a in sys.argv[1:] if not a.startswith("--")]
    if len(pos) < 2:
        sys.exit(__doc__)
    try:
        for c in rendre(pos[0], [int(x) for x in pos[1:]], opt.get("dossier") or dossier_frais(), int(opt.get("dpi", 80))):
            print(c)
    except ValueError as exc:
        print(f"rendu_slide : {exc}", file=sys.stderr)
        sys.exit(2)
