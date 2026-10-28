"""READ-ONLY audit of shape motifs in a .pptx, classified by GEOMETRY (proxy).

Motifs: pastille (numbered / text / plain), filet, bandeau, carte_arrondie, carre. Measures
roundRect corner radii and layout/motif variety. Shapes inside groups are not walked (limit).
Taken over from Vscode8 (scripts/motifs.py) without its slide-mutating helpers
(`harmoniser_rayons`, `arrondir_carres`): mutating helpers stay in Vscode8 until a delta guard
exists. This is an inventory, NOT a charter norm (no charter value exists for corner radii).
Thresholds live in motifs_deck_seuils.json next to this script.

Usage: py motifs_deck.py deck.pptx [--meme-gabarit N]
"""
import collections
import json
import os
import re
import statistics
import sys

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
EMU = 914400.0
SEUILS_DEFAUT = {"pastille_max_in": 0.8, "filet_epaisseur_max_in": 0.12, "filet_longueur_min_in": 0.3,
                 "bandeau_h_min_in": 0.25, "bandeau_h_max_in": 1.2, "bandeau_largeur_ratio": 0.85,
                 "carte_min_w_in": 0.8, "carte_min_h_in": 0.5, "rayon_defaut_adj": 16667,
                 "meme_gabarit_seuil": 0.6}


def charger_seuils(chemin=None):
    """Defaults overridden by the neighbouring JSON (or `chemin`); an unreadable JSON is ignored."""
    seuils = dict(SEUILS_DEFAUT)
    chemin = chemin or os.path.join(os.path.dirname(os.path.abspath(__file__)), "motifs_deck_seuils.json")
    try:
        with open(chemin, encoding="utf-8") as fh:
            lu = json.load(fh)
        if isinstance(lu, dict):
            seuils.update({k: v for k, v in lu.items() if k in SEUILS_DEFAUT})
    except (OSError, ValueError):
        pass
    return seuils


def _prst(sh):
    g = sh._element.spPr.find(A + "prstGeom") if hasattr(sh._element, "spPr") else None
    return g.get("prst") if g is not None else None


def _rempli(sh):
    sp = sh._element.spPr
    return sp.find(A + "solidFill") is not None or sp.find(A + "gradFill") is not None


def classer(sh, larg_slide, seuils=None):
    """Motif name of an autoshape, or None."""
    s = seuils or charger_seuils()
    if sh.shape_type != 1 or not hasattr(sh._element, "spPr"):
        return None
    prst, w, h = _prst(sh), sh.width / EMU, sh.height / EMU
    # txBody check: python-pptx creates an empty txBody on first text_frame access (would mutate)
    texte = sh.text_frame.text.strip() if getattr(sh._element, "txBody", None) is not None else ""
    if prst == "ellipse" and w <= s["pastille_max_in"] and h <= s["pastille_max_in"]:
        return "pastille_numerotee" if re.fullmatch(r"\d{1,2}", texte) else ("pastille_texte" if texte else "pastille")
    if prst in ("rect", "roundRect") and min(w, h) <= s["filet_epaisseur_max_in"] and max(w, h) >= s["filet_longueur_min_in"]:
        return "filet"
    if (prst in ("rect", "roundRect") and w >= s["bandeau_largeur_ratio"] * larg_slide / EMU
            and s["bandeau_h_min_in"] <= h <= s["bandeau_h_max_in"] and _rempli(sh)):
        return "bandeau"
    if prst == "roundRect" and w >= s["carte_min_w_in"] and h >= s["carte_min_h_in"]:
        return "carte_arrondie"
    if prst == "rect" and w >= s["carte_min_w_in"] and h >= s["carte_min_h_in"] and _rempli(sh):
        return "carre"
    return None


def inventaire(prs, seuils=None):
    """{rank: Counter(motif)} per slide."""
    return {i: collections.Counter(m for sh in sl.shapes if (m := classer(sh, prs.slide_width, seuils)))
            for i, sl in enumerate(prs.slides, 1)}


def rayon_in(sh, seuils=None):
    """Corner radius (inches) of a roundRect, or None."""
    s = seuils or charger_seuils()
    if _prst(sh) != "roundRect":
        return None
    gd = sh._element.spPr.find(A + "prstGeom").find(".//" + A + "gd")
    adj = int(gd.get("fmla").split()[1]) if gd is not None else s["rayon_defaut_adj"]
    return adj * min(sh.width, sh.height) / 100000 / EMU


def rayons(prs, seuils=None):
    """Counter of corner radii (rounded to 0.01 in) of every carte_arrondie of the deck."""
    c = collections.Counter()
    for sl in prs.slides:
        for sh in sl.shapes:
            if classer(sh, prs.slide_width, seuils) == "carte_arrondie":
                c[round(rayon_in(sh, seuils), 2)] += 1
    return c


def rayon_dominant(prs, seuils=None):
    """Median radius of the rounded cards (descriptive only), 0.1 if none."""
    vals = [r for r, n in rayons(prs, seuils).items() for _ in range(n)]
    return round(statistics.median(vals), 2) if vals else 0.1


def meme_gabarit(prs, rang_ref, seuils=None):
    """Ranks with the same layout AND a motif signature close to the reference slide."""
    s = seuils or charger_seuils()
    inv = inventaire(prs, s)
    slides = list(prs.slides)
    ref = inv[rang_ref]
    nom = slides[rang_ref - 1].slide_layout.name

    def proche(c):
        cles = set(ref) | set(c)
        if not cles:
            return False
        return sum(min(ref[k], c[k]) for k in cles) / max(sum(max(ref[k], c[k]) for k in cles), 1) >= s["meme_gabarit_seuil"]
    return [i for i, sl in enumerate(slides, 1) if i != rang_ref and sl.slide_layout.name == nom and proche(inv[i])]


def _taille_pt(r):
    """Explicit run size in pt without creating an rPr (python-pptx's `font.size` would mutate)."""
    rpr = r._r.rPr
    sz = rpr.get("sz") if rpr is not None else None
    return int(sz) / 100 if sz else None


def mesures(prs, seuils=None):
    """Variety measures: layouts, distinct motif signatures, bandeaux, text sizes (title / body)."""
    inv = inventaire(prs, seuils)
    sig = collections.Counter(tuple(sorted(c.items())) for c in inv.values())
    tailles, titres = collections.Counter(), collections.Counter()
    for sl in prs.slides:
        for sh in sl.shapes:
            if getattr(sh._element, "txBody", None) is None:  # no text frame; avoids creating one
                continue
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    if _taille_pt(r) and r.text.strip():
                        est_titre = sl.shapes.title is not None and sh.shape_id == sl.shapes.title.shape_id
                        (titres if est_titre else tailles)[_taille_pt(r)] += len(r.text.strip())
    return {"mises_en_page": collections.Counter(sl.slide_layout.name for sl in prs.slides),
            "signatures_distinctes": len(sig), "signature_la_plus_repetee": sig.most_common(1)[0][1] if sig else 0,
            "bandeaux": sum(c["bandeau"] for c in inv.values()),
            "tailles_titres": titres.most_common(5), "tailles_corps": tailles.most_common(8)}


def main(argv=None):
    from pptx import Presentation
    argv = sys.argv[1:] if argv is None else argv
    if not argv:
        print("usage: motifs_deck.py deck.pptx [--meme-gabarit N]", file=sys.stderr)
        return 2
    prs = Presentation(argv[0])
    tot = collections.Counter()
    for c in inventaire(prs).values():
        tot.update(c)
    print("motifs (formes, tout le deck) :", dict(tot))
    print("rayons des cartes arrondies (po -> nombre) :", sorted(rayons(prs).items()))
    print("rayon dominant (mediane) :", rayon_dominant(prs))
    m = mesures(prs)
    print(f"variete : {len(m['mises_en_page'])} mises en page, {m['signatures_distinctes']} signatures de motifs distinctes "
          f"(la plus repetee : {m['signature_la_plus_repetee']} slides), {m['bandeaux']} bandeaux")
    print("tailles des titres (pt, caracteres) :", m["tailles_titres"])
    print("tailles du corps (pt, caracteres) :", m["tailles_corps"])
    if "--meme-gabarit" in argv:
        ref = int(argv[argv.index("--meme-gabarit") + 1])
        print(f"slides de meme gabarit que la slide {ref} : {meme_gabarit(prs, ref)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
