"""pptx_deck_core : operations de reprise de deck, python-pptx seul, sans effet de bord a l'import.

Fonctions NOUVELLES et additives (aucune fonction de pptx_deck.py n'est redefinie) :
  copier_slide(prs, slide, index=None)            copie avec images/liens/notes
  remplacer_texte(prs, ancien, nouveau, ...)      remplacement qui garde les runs
  renumeroter_slides(prs, debut=1, ...)           champs/numeros de slide
  extract_comments(chemin)                        commentaires classiques + modernes
  diff_decks(base, retour, slides=None)           slides modifiees et nature du changement
Ajouts "delta" (portes de Vscode8 deck_lib, convention du core : selecteur titre/id/rang) :
  geometrie_forme(prs, slide, x,y,w,h, nom=, texte=)  edite la geometrie d'une forme
  mesure_forme(slide, shape_id)                   estimation de debordement d'une zone de texte
  deplacer_slide(prs, selecteur, position)        reordonne, sections p14 preservees
  garder_slides(prs, selecteurs)                  ne garde que ces slides, sections nettoyees
  couleurs_en_theme(prs)                          srgbClr egal a un accent -> schemeClr
  empreintes_slides / run_delta(src, dst, edits)  delta complet : refuse d'ecrire dst si une
                                                  slide hors perimetre a change, dst != src
Filtre commun `slides` : iterable de rangs 1-based (None = toutes). `dry_run=True` : liste
les slides touchees sans rien modifier. Convention des commentaires alignee sur
Vscode8 feedbacks_reprise.commentaires (rang, auteur, texte), rendue ici en dicts.
"""
import copy
import os
import re
import sys
import zipfile

from lxml import etree

_NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
_NS_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_SAUTE = ("slideLayout", "notesSlide")
_FORMES = ("sp", "pic", "graphicFrame", "grpSp", "cxnSp")


def _ouvrir(deck):
    from pptx import Presentation
    return Presentation(deck) if isinstance(deck, str) or hasattr(deck, "read") else deck


def _rangs(n, slides):
    return set(range(1, n + 1)) if slides is None else {int(s) for s in slides}


def id_slide(slide):
    """Identifiant `sldId` (stable quand les slides bougent) d'une slide."""
    return slide.slide_id


def resoudre_slides(prs, slides):
    """Selecteurs -> rangs 1-based (ensemble). Un selecteur est : un entier (rang), une
    chaine (TITRE : egalite stricte normalisee sur le texte d'une forme), `"id:256"` ou
    `{"id": 256}` (sldId). None = toutes. Titre ambigu ou introuvable : ValueError qui
    liste les candidats ; rang hors deck : ValueError."""
    n = len(prs.slides)
    if slides is None:
        return set(range(1, n + 1))
    if isinstance(slides, (int, str, dict)):
        slides = [slides]
    ids = {sl.slide_id: i for i, sl in enumerate(prs.slides, 1)}
    sortie = set()
    for sel in slides:
        if isinstance(sel, dict) or (isinstance(sel, str) and sel.startswith("id:")):
            sid = int(sel["id"] if isinstance(sel, dict) else sel[3:])
            if sid not in ids:
                raise ValueError(f"sldId {sid} introuvable ; ids presents : {sorted(ids)}")
            sortie.add(ids[sid])
        elif isinstance(sel, str):
            cible = _norm(sel)
            trouvees = [i for i, sl in enumerate(prs.slides, 1)
                        if any(sh.has_text_frame and _norm(sh.text_frame.text) == cible
                               for sh in sl.shapes)]
            if len(trouvees) != 1:
                raise ValueError(f"titre {sel!r} : {len(trouvees)} slide(s) trouvee(s) "
                                 f"({', '.join(map(str, trouvees)) or 'aucune'}) ; 1 exigee")
            sortie.add(trouvees[0])
        else:
            if not 1 <= int(sel) <= n:
                raise ValueError(f"rang {sel} hors deck (1..{n})")
            sortie.add(int(sel))
    return sortie


def _norm(t):
    return re.sub(r"\s+", " ", (t or "").replace(" ", " ")).strip().casefold()


def _local(e):
    return e.tag.split("}")[-1]


def _cloner_partie(partie, paquet):
    """Clone une partie (graphique, classeur embarque) sous un nouveau partname, avec des
    relations propres : editer la copie ne touche jamais l'original."""
    from pptx.opc.package import Part
    from pptx.opc.packuri import PackURI
    modele = re.sub(r"\d+(\.\w+)$", r"%d\1", str(partie.partname))
    if "%d" not in modele:
        modele = re.sub(r"(\.\w+)$", r"%d\1", str(partie.partname))
    neuve = Part(PackURI(paquet.next_partname(modele)), partie.content_type, paquet, partie.blob)
    for rel in partie.rels.values():
        if rel.is_external:
            neuve.rels.get_or_add_ext_rel(rel.reltype, rel.target_ref)
        elif "package" in rel.reltype:
            neuve.relate_to(_cloner_partie(rel.target_part, paquet), rel.reltype)
        else:
            neuve.relate_to(rel.target_part, rel.reltype)
    # le blob d'un graphique cite ses relations par rId : reconstruire dans le meme ordre
    return neuve


_PARTAGEABLES = ("/image", "/hyperlink", "/slide")  # immuables ou simples renvois


def copier_slide(prs, slide, index=None):
    """Copie `slide` (meme layout) avec ses relations ; renvoie la copie.

    Images et liens : partages (immuables). Notes : texte recopie. Graphiques : partie
    ET classeur embarque clones (la copie est independante). Toute autre partie liee
    (media, objet OLE, diagramme...) leve ValueError : mieux vaut un refus qu'une copie qui
    partage silencieusement une partie avec l'original.
    Limites : une slide a effets d'image Office 2007 (relation hdphoto) est refusee ; le
    graphique cloné est une partie generique : enregistrer puis rouvrir le deck avant
    d'utiliser `shape.chart` sur la copie.
    `index` : position 0-based dans le deck (defaut : a la fin)."""
    a_cloner, a_partager = {}, {}
    for rid, rel in slide.part.rels.items():
        if any(k in rel.reltype for k in _SAUTE):
            continue
        if rel.is_external or rel.reltype.endswith(_PARTAGEABLES):
            a_partager[rid] = rel
        elif rel.reltype.endswith("/chart"):
            a_cloner[rid] = rel
        else:
            raise ValueError(
                f"copier_slide : partie liee non clonable ({rel.reltype}, rId {rid}) ; "
                "copie refusee plutot que partagee avec l'original")
    neuve = prs.slides.add_slide(slide.slide_layout)
    for el in list(neuve.shapes._spTree):
        if _local(el) in _FORMES:
            neuve.shapes._spTree.remove(el)
    rid_map = {}
    for rid, rel in a_partager.items():
        if rel.is_external:
            rid_map[rid] = neuve.part.relate_to(rel.target_ref, rel.reltype, is_external=True)
        else:
            rid_map[rid] = neuve.part.relate_to(rel.target_part, rel.reltype)
    for rid, rel in a_cloner.items():
        rid_map[rid] = neuve.part.relate_to(
            _cloner_partie(rel.target_part, prs.part.package), rel.reltype)
    for el in slide.shapes._spTree:
        if _local(el) in _FORMES:
            c = copy.deepcopy(el)
            for n in c.iter():
                for k, v in list(n.attrib.items()):
                    if k.startswith(f"{{{_NS_R}}}") and v in rid_map:
                        n.set(k, rid_map[v])
            neuve.shapes._spTree.append(c)
    if slide.has_notes_slide:
        neuve.notes_slide.notes_text_frame.text = slide.notes_slide.notes_text_frame.text
    if index is not None:
        lst = prs.slides._sldIdLst
        el = lst[-1]
        lst.remove(el)
        lst.insert(index, el)
    return neuve


def _paragraphes(shape):
    if shape.has_text_frame:
        yield from shape.text_frame.paragraphs
    if getattr(shape, "has_table", False) and shape.has_table:
        for r in shape.table.rows:
            for c in r.cells:
                yield from c.text_frame.paragraphs
    if shape.shape_type == 6:  # groupe
        for s in shape.shapes:
            yield from _paragraphes(s)


def _remplacer_para(p, ancien, nouveau):
    """Remplace chaque occurrence ; le texte neuf prend la mise en forme du 1er run touche."""
    n, depart = 0, 0
    while ancien:
        runs = list(p.runs)
        txt = "".join(r.text for r in runs)
        i = txt.find(ancien, depart)
        if i < 0:
            return n
        j, pos, premier = i + len(ancien), 0, True
        for r in runs:
            a, b = pos, pos + len(r.text)
            pos = b
            if b <= i or a >= j:
                continue
            avant = r.text[:max(i - a, 0)]
            apres = r.text[j - a:] if j < b else ""
            r.text = avant + (nouveau if premier else "") + apres
            premier = False
        n += 1
        depart = i + len(nouveau)
    return n


def remplacer_texte(prs, ancien, nouveau, slides=None, dry_run=False):
    """Remplace `ancien` par `nouveau` en gardant les runs (gras, couleur du 1er run de
    l'occurrence). Renvoie {rang_slide: nb_occurrences}. `dry_run` : ne modifie rien."""
    sortie = {}
    for rang in sorted(resoudre_slides(prs, slides)):
        if rang < 1 or rang > len(prs.slides):
            continue
        n = 0
        for shp in prs.slides[rang - 1].shapes:
            for p in _paragraphes(shp):
                if dry_run:
                    n += "".join(r.text for r in p.runs).count(ancien) if ancien else 0
                else:
                    n += _remplacer_para(p, ancien, nouveau)
        if n:
            sortie[rang] = n
    return sortie


def renumeroter_slides(prs, debut=1, slides=None, dry_run=False):
    """Met a jour les numeros affiches (champs `slidenum`). Renvoie {rang: numero}.
    `dry_run` : liste sans modifier."""
    sortie = {}
    for rang in sorted(resoudre_slides(prs, slides)):
        if rang < 1 or rang > len(prs.slides):
            continue
        num = str(debut + rang - 1)
        for el in prs.slides[rang - 1].shapes._spTree.iter(f"{{{_NS_A}}}fld"):
            if (el.get("type") or "").startswith("slidenum"):
                t = el.find(f"{{{_NS_A}}}t")
                if t is not None:
                    sortie[rang] = int(num)
                    if not dry_run:
                        t.text = num
    return sortie


def extract_comments(chemin):
    """[{slide, auteur, texte, forme}] : commentaires classiques (comment*.xml) et modernes
    (modernComment_*.xml). `forme` = id de forme ciblee (moderne, spMk) ou None."""
    prs = _ouvrir(chemin)
    z = zipfile.ZipFile(chemin)
    auteurs = {}
    for n in z.namelist():
        if n.endswith("authors.xml") or n.endswith("commentAuthors.xml"):
            for e in etree.fromstring(z.read(n)).iter():
                if _local(e) in ("author", "cmAuthor") and e.get("id") is not None:
                    auteurs[e.get("id")] = e.get("name", "")
    sortie = []
    for i, s in enumerate(prs.slides, 1):
        for rel in s.part.rels.values():
            if "comment" not in rel.reltype.lower() or rel.is_external:
                continue
            x = etree.fromstring(z.read(str(rel.target_part.partname).lstrip("/")))
            for cm in x.iter():
                if _local(cm) != "cm":
                    continue
                paras = ["".join(t.text or "" for t in p.iter() if _local(t) == "t")
                         for p in cm.iter() if _local(p) == "p"]
                txt = " | ".join(p for p in paras if p) or "".join(
                    t.text or "" for t in cm.iter() if _local(t) == "text")
                forme = next((e.get("id") for e in cm.iter() if _local(e) == "spMk"), None)
                aid = cm.get("authorId", "")
                sortie.append({"slide": i, "auteur": auteurs.get(aid, aid or "?"),
                               "texte": txt.strip(), "forme": forme})
    return sortie


def _empreinte(slide):
    out = {}
    for shp in slide.shapes:
        geo = (shp.left, shp.top, shp.width, shp.height)
        out[shp.shape_id] = (geo, "\n".join(p.text for p in _paragraphes(shp)))
    return out


def _cle_texte(slide):
    return chr(31).join(t for t in (chr(10).join(p.text for p in _paragraphes(s)) for s in slide.shapes) if t)


def _natures(sa, sb):
    ea, eb = _empreinte(sa), _empreinte(sb)
    nat = []
    if set(eb) - set(ea):
        nat.append("forme_ajoutee")
    if set(ea) - set(eb):
        nat.append("forme_supprimee")
    com = set(ea) & set(eb)
    if any(ea[k][1] != eb[k][1] for k in com):
        nat.append("texte")
    if any(ea[k][0] != eb[k][0] for k in com):
        nat.append("geometrie")
    return nat


def diff_decks(base, retour, slides=None):
    """Compare deux decks en APPARIANT les slides (empreinte du texte, puis blocs de meme
    taille) avant de comparer : une insertion ou suppression est signalee comme telle et ne
    decale pas les autres. {rang: [natures]} parmi 'texte', 'geometrie', 'forme_ajoutee',
    'forme_supprimee', 'slide_ajoutee' (rang du retour), 'slide_supprimee' (rang de la base).
    `slides` filtre sur ces rangs. Seules les slides modifiees figurent au resultat."""
    import difflib
    a, b = _ouvrir(base), _ouvrir(retour)
    la, lb = list(a.slides), list(b.slides)
    ka, kb = [_cle_texte(s) for s in la], [_cle_texte(s) for s in lb]
    sortie = {}
    for op, i1, i2, j1, j2 in difflib.SequenceMatcher(None, ka, kb, autojunk=False).get_opcodes():
        if op == "equal" or op == "replace":
            n = min(i2 - i1, j2 - j1)
            for k in range(n):
                nat = _natures(la[i1 + k], lb[j1 + k])
                if nat:
                    sortie[j1 + k + 1] = nat
            for k in range(n, j2 - j1):
                sortie[j1 + k + 1] = ["slide_ajoutee"]
            for k in range(n, i2 - i1):
                sortie[i1 + k + 1] = sortie.get(i1 + k + 1, []) + ["slide_supprimee"]
        elif op == "insert":
            for k in range(j1, j2):
                sortie[k + 1] = ["slide_ajoutee"]
        elif op == "delete":
            for k in range(i1, i2):
                sortie[k + 1] = sortie.get(k + 1, []) + ["slide_supprimee"]
    cible = _rangs(max(len(la), len(lb)), slides)
    return {r: v for r, v in sorted(sortie.items()) if r in cible}


def trouver_forme(slide, nom=None, texte=None):
    """LA forme de `slide` dont le nom egale `nom` et/ou dont le texte contient `texte`.
    ValueError explicite (candidats listes) si zero ou plusieurs formes."""
    if nom is None and texte is None:
        raise ValueError("trouver_forme : nom ou texte exige")
    trouvees = [sh for sh in slide.shapes
                if (nom is None or sh.name == nom)
                and (texte is None or (sh.has_text_frame and texte in sh.text_frame.text))]
    if len(trouvees) != 1:
        cand = ", ".join(f"{sh.name!r}" for sh in slide.shapes)
        raise ValueError(f"trouver_forme(nom={nom!r}, texte={texte!r}) : {len(trouvees)} "
                         f"forme(s) ; formes de la slide : {cand}")
    return trouvees[0]


_P14 = "http://schemas.microsoft.com/office/powerpoint/2010/main"
_URI_SECTIONS = "{521415D9-36F7-43E2-AB2F-B90AF26B5E84}"
_NS_P = "http://schemas.openxmlformats.org/presentationml/2006/main"


def definir_sections(prs, sections):
    """Ecrit (ou remplace) p14:sectionLst. `sections` : [(nom, [selecteurs de slides])] ;
    les slides non citees sont ajoutees a la derniere section. Les ids sont ceux des sldId."""
    import uuid
    ordre = [sl.slide_id for sl in prs.slides]
    pris, blocs = set(), []
    for nom, sels in sections:
        ids = [ordre[r - 1] for r in sorted(resoudre_slides(prs, sels))]
        doublons = pris & set(ids)
        if doublons:
            raise ValueError(f"slides dans deux sections : {sorted(doublons)}")
        pris |= set(ids)
        blocs.append((nom, ids))
    reste = [i for i in ordre if i not in pris]
    if reste and blocs:
        blocs[-1] = (blocs[-1][0], blocs[-1][1] + reste)
    racine = prs.part._element
    ext_lst = racine.find(f"{{{_NS_P}}}extLst")
    if ext_lst is None:
        ext_lst = etree.SubElement(racine, f"{{{_NS_P}}}extLst")
    for ext in list(ext_lst):
        if ext.get("uri") == _URI_SECTIONS:
            ext_lst.remove(ext)
    ext = etree.SubElement(ext_lst, f"{{{_NS_P}}}ext", uri=_URI_SECTIONS)
    lst = etree.SubElement(ext, f"{{{_P14}}}sectionLst", nsmap={"p14": _P14})
    for nom, ids in blocs:
        sec = etree.SubElement(lst, f"{{{_P14}}}section", name=nom,
                               id="{" + str(uuid.uuid4()).upper() + "}")
        sl = etree.SubElement(sec, f"{{{_P14}}}sldIdLst")
        for i in ids:
            etree.SubElement(sl, f"{{{_P14}}}sldId", id=str(i))
    return blocs


def supprimer_slide_par_selecteur(prs, selecteur):
    """Retire une slide du deck (relation lachee : la partie n'est plus ecrite)."""
    (rang,) = resoudre_slides(prs, [selecteur])
    _retirer_slide(prs, prs.slides._sldIdLst[rang - 1])


def _retirer_slide(prs, el):
    """Retire le sldId `el`, lache sa relation et le retire d'un sectionLst existant
    (sinon PowerPoint peut proposer une reparation a l'ouverture)."""
    ident = str(el.id)
    prs.part.drop_rel(el.rId)
    prs.slides._sldIdLst.remove(el)
    for sid in prs.part._element.iter(f"{{{_P14}}}sldId"):
        if sid.get("id") == ident:
            sid.getparent().remove(sid)


_OPS = ("remplacer_texte", "supprimer", "copier", "renumeroter", "deplacer", "geometrie", "garder")


def appliquer_delta(prs, delta, dry_run=False):
    """Delta declaratif : liste de {slide: <titre|id|rang>, op: ..., ...} avec op parmi
    remplacer_texte (ancien, nouveau), supprimer, copier (index optionnel), renumeroter
    (debut), deplacer (position 1-based, sections preservees), geometrie (forme: {nom|texte,
    x, y, w, h en EMU, None = inchange}) et garder (slides: [selecteurs], sans cle `slide` :
    ne garde que celles-la). Les selecteurs sont resolus AVANT toute modification (une
    insertion ou suppression ne decale rien) ; ambigu/introuvable : ValueError. Renvoie
    [{slide: rang_initial, op}] ; `dry_run` : liste sans rien modifier."""
    ops = []
    for d in delta:
        if d.get("op") not in _OPS:
            raise ValueError(f"op inconnue : {d.get('op')!r}")
        if d["op"] == "garder":
            if not d.get("slides"):
                raise ValueError("garder : liste de slides vide refusee")
            ops.append((None, dict(d, _gardes=sorted(resoudre_slides(prs, d["slides"])))))
            continue
        sel = d.get("slide")
        if sel is None and d["op"] in ("deplacer", "geometrie", "supprimer", "copier"):
            raise ValueError(f"{d['op']} : cle `slide` exigee")
        (rang,) = resoudre_slides(prs, [sel]) if sel is not None else (None,)
        ops.append((rang, d))
    cibles = {r: prs.slides[r - 1] for r, _ in ops if r is not None}
    toutes = list(prs.slides)
    rapport = []
    for rang, d in ops:
        op = d["op"]
        if op == "remplacer_texte":
            hit = remplacer_texte(prs, d["ancien"], d["nouveau"],
                                  slides=None if rang is None else [id_ref(cibles[rang], prs)],
                                  dry_run=dry_run)
            for r in hit:
                rapport.append({"slide": r, "op": op})
            continue
        if not dry_run:
            if op == "supprimer":
                el = next(e for e in prs.slides._sldIdLst if e.id == cibles[rang].slide_id)
                _retirer_slide(prs, el)
            elif op == "copier":
                copier_slide(prs, cibles[rang], index=d.get("index"))
            elif op == "renumeroter":
                renumeroter_slides(prs, debut=d.get("debut", 1))
            elif op == "deplacer":
                deplacer_slide(prs, id_ref(cibles[rang], prs), d["position"])
            elif op == "geometrie":
                f = d["forme"]
                geometrie_forme(prs, id_ref(cibles[rang], prs), f.get("x"), f.get("y"),
                                f.get("w"), f.get("h"), nom=f.get("nom"), texte=f.get("texte"))
            elif op == "garder":
                garder_slides(prs, [id_ref(toutes[r - 1], prs) for r in d["_gardes"]])
        rapport.append({"slide": rang, "op": op})
    return rapport


def geometrie_forme(prs, slide, x=None, y=None, w=None, h=None, nom=None, texte=None):
    """Edite la geometrie (EMU ; None = inchange) de LA forme (`trouver_forme`) de la slide
    designee par selecteur. ValueError si slide/forme ambigue ou introuvable. Renvoie la forme."""
    (rang,) = resoudre_slides(prs, [slide])
    sh = trouver_forme(prs.slides[rang - 1], nom=nom, texte=texte)
    for attr, v in (("left", x), ("top", y), ("width", w), ("height", h)):
        if v is not None:
            setattr(sh, attr, int(v))
    return sh


def mesure_forme(slide, shape_id):
    """Estimation de debordement d'une zone de texte : {chars, taille_pt, ratio} ou None.
    ratio = hauteur de texte estimee / hauteur utile (glyphe ~ 0,5 x taille) ; > 1 : debordement
    probable ; taille None = heritee. Estimation grossiere : marge large (0,9 a 1,15 = zone grise)."""
    import math
    sh = next((x for x in slide.shapes if x.shape_id == shape_id), None)
    if sh is None or not sh.has_text_frame or not sh.width or not sh.height:
        return None
    tailles = [r.font.size.pt for p in sh.text_frame.paragraphs for r in p.runs if r.font.size]
    taille = max(set(tailles), key=tailles.count) if tailles else None
    t = taille or 18.0
    largeur = max(sh.width / 12700 - 14.4, 1.0)
    lignes = sum(max(1, math.ceil(len(p.text) * 0.5 * t / largeur))
                 for p in sh.text_frame.paragraphs)
    return {"chars": len(_norm(sh.text_frame.text)), "taille_pt": taille,
            "ratio": round(lignes * t * 1.2 / max(sh.height / 12700 - 7.2, 1.0), 2)}


def _appartenance(prs):
    """{sldId(str): rang de section p14} avant une operation de structure."""
    secs = prs.part._element.findall(f".//{{{_P14}}}section")
    return {e.get("id"): i for i, sec in enumerate(secs) for e in sec.iter(f"{{{_P14}}}sldId")}


def _reconstruire_sections(prs, appart):
    """Recompose les sldIdLst des sections dans l'ordre reel ; slide inconnue : section de la
    precedente (la 1re section si elle est en tete). Sans sectionLst : no-op."""
    secs = prs.part._element.findall(f".//{{{_P14}}}section")
    if not secs:
        return
    prec, membre = 0, {}
    ordre = [str(e.id) for e in prs.slides._sldIdLst]
    for sid in ordre:
        prec = appart.get(sid, prec)
        membre[sid] = prec
    for i, sec in enumerate(secs):
        lst = sec.find(f"{{{_P14}}}sldIdLst")
        for e in list(lst):
            lst.remove(e)
        for sid in ordre:
            if membre[sid] == i:
                etree.SubElement(lst, f"{{{_P14}}}sldId").set("id", sid)


def deplacer_slide(prs, selecteur, position):
    """Place la slide au rang `position` (1-based, 1 = premiere) ; sa section devient celle de
    la slide qui la precede. ValueError si position hors 1..n."""
    n = len(prs.slides)
    if isinstance(position, bool) or not isinstance(position, int) or not 1 <= position <= n:
        raise ValueError(f"position {position!r} invalide (entier de 1 a {n})")
    (rang,) = resoudre_slides(prs, [selecteur])
    appart = _appartenance(prs)
    lst = prs.slides._sldIdLst
    el = lst[rang - 1]
    lst.remove(el)
    lst.insert(position - 1, el)
    appart.pop(str(el.id), None)
    _reconstruire_sections(prs, appart)


def garder_slides(prs, selecteurs):
    """Ne garde que les slides designees (les autres : relation lachee, sections nettoyees).
    Liste vide refusee (ValueError) : un deck sans slide n'est pas un resultat voulu."""
    if not selecteurs:
        raise ValueError("garder_slides : aucune slide a garder")
    gardes = resoudre_slides(prs, selecteurs)
    lst = prs.slides._sldIdLst
    for i, el in reversed(list(enumerate(list(lst), 1))):
        if i not in gardes:
            _retirer_slide(prs, el)


def couleurs_en_theme(prs):
    """Remplace dans les slides chaque srgbClr egale a un accent du theme (accent1..6) par le
    schemeClr correspondant (enfants alpha/lumMod gardes). Renvoie le nombre de remplacements."""
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    theme = etree.fromstring(prs.slide_master.part.part_related_by(RT.THEME).blob)
    carte = {}
    for k in range(1, 7):
        el = theme.find(f".//{{{_NS_A}}}accent{k}/{{{_NS_A}}}srgbClr")
        if el is not None:
            carte.setdefault(el.get("val").upper(), f"accent{k}")
    n = 0
    for sl in prs.slides:
        for el in sl._element.iter(f"{{{_NS_A}}}srgbClr"):
            nom = carte.get((el.get("val") or "").upper())
            if nom:
                el.tag = f"{{{_NS_A}}}schemeClr"
                el.set("val", nom)
                n += 1
    return n


def empreintes_slides(prs):
    """{sldId: (empreinte texte+geometrie+XML des formes, nb formes)} : stable quand les slides
    bougent ; l'XML des formes rend visible un changement de mise en forme (gras, taille)."""
    import hashlib
    sortie = {}
    for sl in prs.slides:
        emp = repr(sorted(_empreinte(sl).items())) + etree.tostring(sl.shapes._spTree).decode("utf8")
        sortie[sl.slide_id] = (hashlib.md5(emp.encode("utf8")).hexdigest(), len(sl.shapes))
    return sortie


def _integrite():
    """Import paresseux de integrite_pptx (skill pptx-verify) ; absent = ValueError (un delta ne sort
    jamais sans son controle OPC). Cherche a cote (kit : scripts/../../pptx-verify/scripts) et depuis
    le canon du hub (dispositif/canon/../../skills/pptx-verify/scripts)."""
    ici = os.path.dirname(os.path.abspath(__file__))
    for rel in (("..", "..", "pptx-verify", "scripts"), ("..", "..", "skills", "pptx-verify", "scripts")):
        d = os.path.normpath(os.path.join(ici, *rel))
        if os.path.isfile(os.path.join(d, "integrite_pptx.py")):
            if d not in sys.path:
                sys.path.append(d)
            import integrite_pptx
            return integrite_pptx
    raise ValueError("run_delta : integrite_pptx introuvable ; propager la skill pptx-verify (kit agentic) dans ce projet")


def run_delta(src, dst, edits, dry=False):
    """Ouvre `src`, applique `edits` (cf. appliquer_delta), compare avant/apres PAR sldId et
    REFUSE d'ecrire `dst` (ValueError) si une slide hors perimetre a change. Perimetre = slides
    visees par une op avec `slide` ; une op sans `slide` (remplacer_texte global, renumeroter)
    couvre toutes les slides, a condition qu'aucune op ciblee ne soit melangee (ValueError). `dst` doit differer de `src` (meme fichier : ValueError, rien
    d'ecrit). `dry` : controle les selecteurs sans modifier ni ecrire. Renvoie le rapport."""
    s_abs, d_abs = os.path.abspath(src), os.path.abspath(dst)
    if os.path.normcase(s_abs) == os.path.normcase(d_abs) or (
            os.path.exists(d_abs) and os.path.samefile(s_abs, d_abs)):
        raise ValueError("run_delta : dst == src refuse (ne jamais ecraser la source)")
    prs, avant = _ouvrir(src), _ouvrir(src)
    emp_avant = empreintes_slides(avant)
    portee, tout = set(), False
    for d in edits:
        if d.get("op") == "garder":
            continue
        if d.get("slide") is None:
            tout = True
            continue
        (rang,) = resoudre_slides(prs, [d["slide"]])
        portee.add(prs.slides[rang - 1].slide_id)
    if tout and portee:
        raise ValueError("run_delta : melange d'ops ciblees et globales refuse (le controle des "
                         "slides figees serait desactive) ; separer en deux appels")
    rapport = appliquer_delta(prs, edits, dry_run=dry)
    if not dry and not tout:
        apres = empreintes_slides(prs)
        hors = [i for i in apres if i in emp_avant and apres[i] != emp_avant[i] and i not in portee]
        if hors:
            raise ValueError(f"run_delta : slide(s) hors perimetre modifiee(s) (sldId {hors}) ; "
                             "dst non ecrit")
    if not dry:
        tmp = dst + ".tmp"
        try:
            prs.save(tmp)
            pb = _integrite().controler(tmp)["problemes"]
            if pb:
                raise ValueError("run_delta : integrite OPC KO : " + " ; ".join(pb) + " ; dst inchange")
            os.replace(tmp, dst)
        finally:
            if os.path.exists(tmp):
                os.remove(tmp)
    return rapport


def id_ref(slide, prs):
    """Selecteur `id:N` d'une slide (resolution stable apres insertion/suppression)."""
    return f"id:{slide.slide_id}"
