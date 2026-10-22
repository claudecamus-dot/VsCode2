"""Recettes de slides « data seule » : l'appelant donne du CONTENU, jamais de coordonnees.

    import pptx_recettes as R
    print(R.lister())                         # 1 listing court au lieu du catalogue de patterns
    R.recette(slide, "kpi_row", {"kpis": [...]}, zone=None)

`zone` = (x, y, w, h) en EMU ; par defaut, la zone libre sous le titre de la slide
(marges 0,5 in, bas reserve 0,6 in). Chaque recette :
  1. valide `data` contre le tableau SPEC (nombre d'elements, longueurs maximales,
     cles inconnues refusees) -> ValueError au message precis ;
  2. prend couleurs et police dans le THEME du deck (theme_colors / police_theme),
     jamais en hexadecimal code en dur ;
  3. choisit UNE taille de police par bloc via ajuster_police, entre TYPE["body"] et
     TYPE["tiny"] ; si le texte ne tient pas a la taille plancher -> ValueError
     (pas de pagination : couper en deux slides est la decision de l'appelant) ;
  4. finit par verifier_geometrie + verifier_debordements_texte sur CETTE slide et
     verifie que chaque forme posee reste dans `zone` -> RecetteDefaut sinon.

Rendu : la geometrie et le debordement sont verifies par estimation, pas au rendu.
Qualite visuelle non verifiee au rendu tant qu'un deck reel n'a pas ete ouvert en image.
"""
import copy
import os
import re
import sys

_ICI = os.path.dirname(os.path.abspath(__file__))
for _p in (_ICI, os.path.join(_ICI, "..", "..", "skills", "pptx-deck", "scripts")):
    if os.path.isdir(_p) and _p not in sys.path:
        sys.path.append(_p)

import pptx_deck as D  # noqa: E402
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN  # noqa: E402
from pptx.util import Emu  # noqa: E402

GAP = 0.2          # pouces entre blocs
MARGE_H = 0.5      # marge laterale / haute par defaut
BAS_RESERVE = 0.6  # bas de slide laisse au numero de page
CPI_PESSIMISTE = 10.7  # meme calibration que verifier_debordements_texte
T_MAX = D.TYPE["body"]
T_MIN = D.TYPE["tiny"]


class RecetteDefaut(RuntimeError):
    """La recette a produit une slide qui echoue aux controles geometrie/debordement."""


# --- Tableau des limites (source unique : validation ET lister()) -----------------
# champ -> (type, longueur max, requis[, (min, max) pour une liste de textes])
SPEC = {
    "kpi_row": {
        "top": {},
        "liste": ("kpis", (3, 5), {"value": ("s", 8, True), "label": ("s", 40, True),
                                   "note": ("s", 60, False)}),
        "doc": "{kpis: [{value<=8, label<=40, note<=60 opt}]} 3-5 kpis",
    },
    "comparison": {
        "top": {},
        "liste": ("columns", (2, 3), {"title": ("s", 30, True),
                                      "points": ("ls", 90, True, (2, 6))}),
        "doc": "{columns: [{title<=30, points: [2-6 textes <=90]}]} 2-3 colonnes",
    },
    "timeline": {
        "top": {},
        "liste": ("steps", (3, 7), {"date": ("s", 12, True), "title": ("s", 30, True),
                                    "text": ("s", 80, False)}),
        "doc": "{steps: [{date<=12, title<=30, text<=80 opt}]} 3-7 etapes",
    },
    "process": {
        "top": {},
        "liste": ("steps", (3, 6), {"title": ("s", 24, True), "text": ("s", 70, True)}),
        "doc": "{steps: [{title<=24, text<=70}]} 3-6 etapes chevron",
    },
    "cards_grid": {
        "top": {},
        "liste": ("cards", (4, 6), {"title": ("s", 40, True), "text": ("s", 160, True)}),
        "doc": "{cards: [{title<=40, text<=160}]} exactement 4 (2x2) ou 6 (3x2)",
    },
    "quote": {
        "top": {"quote": ("s", 220, True), "author": ("s", 60, True),
                "role": ("s", 60, False)},
        "liste": None,
        "doc": "{quote<=220, author<=60, role<=60 opt}",
    },
    "highlight": {
        "top": {"value": ("s", 10, True), "label": ("s", 60, True),
                "points": ("ls", 100, True, (2, 4))},
        "liste": None,
        "doc": "{value<=10, label<=60, points: [2-4 textes <=100]} chiffre cle + arguments",
    },
    "matrix_2x2": {
        "top": {"x_axis": ("s", 30, True), "y_axis": ("s", 30, True)},
        "liste": ("quadrants", (4, 4), {"title": ("s", 30, True), "text": ("s", 110, True)}),
        "doc": "{x_axis<=30, y_axis<=30, quadrants: [4 x {title<=30, text<=110}]} "
               "ordre : haut-gauche, haut-droite, bas-gauche, bas-droite",
    },
}

EXEMPLES = {
    "kpi_row": {"kpis": [{"value": "87 %", "label": "Taux d'adoption", "note": "T3 2026"},
                         {"value": "12", "label": "Equipes"},
                         {"value": "3,4 M", "label": "Economie annuelle"}]},
    "comparison": {"columns": [
        {"title": "Option A", "points": ["Cout initial faible", "Deploiement en 2 semaines"]},
        {"title": "Option B", "points": ["Cout initial eleve", "Evolutif a long terme",
                                         "Support editeur inclus"]}]},
    "timeline": {"steps": [{"date": "Jan 2026", "title": "Cadrage", "text": "Perimetre valide"},
                           {"date": "Avr 2026", "title": "Pilote"},
                           {"date": "Sept 2026", "title": "Generalisation",
                            "text": "Toutes les equipes"}]},
    "process": {"steps": [{"title": "Collecter", "text": "Exporter les donnees brutes"},
                          {"title": "Analyser", "text": "Mesurer les ecarts"},
                          {"title": "Decider", "text": "Arbitrer les actions"}]},
    "cards_grid": {"cards": [{"title": f"Axe {i}", "text": "Un constat court et factuel."}
                             for i in range(1, 5)]},
    "quote": {"quote": "La qualite se mesure a ce que l'on peut verifier.",
              "author": "A. Martin", "role": "Directrice technique"},
    "highlight": {"value": "-40 %", "label": "Temps de livraison",
                  "points": ["Automatisation des tests", "Revue en parallele"]},
    "matrix_2x2": {"x_axis": "Effort", "y_axis": "Impact", "quadrants": [
        {"title": "Quick wins", "text": "A lancer tout de suite"},
        {"title": "Projets", "text": "A planifier"},
        {"title": "Bonus", "text": "Si le temps le permet"},
        {"title": "A eviter", "text": "Cout superieur au gain"}]},
}


# --- Validation ------------------------------------------------------------------
def _champ(valeur, spec, chemin):
    kind, maxlen, requis = spec[0], spec[1], spec[2]
    if valeur is None:
        if requis:
            raise ValueError(f"{chemin}: champ requis absent")
        return
    if kind == "ls":
        mini, maxi = spec[3]
        if not isinstance(valeur, list) or not mini <= len(valeur) <= maxi:
            n = len(valeur) if isinstance(valeur, list) else type(valeur).__name__
            raise ValueError(f"{chemin}: attendu une liste de {mini} a {maxi} textes, recu {n}")
        for i, v in enumerate(valeur):
            _champ(v, ("s", maxlen, True), f"{chemin}[{i}]")
        return
    if not isinstance(valeur, str) or not valeur.strip():
        raise ValueError(f"{chemin}: texte non vide attendu")
    if len(valeur) > maxlen:
        raise ValueError(f"{chemin}: {len(valeur)} caracteres, max {maxlen}")


def _champs(d, champs, chemin):
    if not isinstance(d, dict):
        raise ValueError(f"{chemin}: objet attendu, recu {type(d).__name__}")
    inconnus = sorted(set(d) - set(champs))
    if inconnus:
        raise ValueError(f"{chemin}: cles inconnues {inconnus} (attendu {sorted(champs)})")
    for nom, spec in champs.items():
        _champ(d.get(nom), spec, f"{chemin}.{nom}")


def valider(nom, data):
    """Valide `data` contre SPEC[nom] ; ValueError precise, sinon renvoie data."""
    spec = SPEC[nom]
    cle_liste = spec["liste"][0] if spec["liste"] else None
    if not isinstance(data, dict):
        raise ValueError(f"{nom}: data doit etre un objet, recu {type(data).__name__}")
    top = {k: v for k, v in data.items() if k != cle_liste}
    _champs(top, spec["top"], nom)
    if cle_liste:
        (_, (mini, maxi), champs) = spec["liste"]
        items = data.get(cle_liste)
        if not isinstance(items, list) or not mini <= len(items) <= maxi:
            n = len(items) if isinstance(items, list) else "absent"
            attendu = f"exactement {mini}" if mini == maxi else f"entre {mini} et {maxi}"
            if nom == "cards_grid":
                attendu = "exactement 4 ou 6"
            raise ValueError(f"{nom}.{cle_liste}: {attendu} elements attendus, recu {n}")
        if nom == "cards_grid" and len(items) not in (4, 6):
            raise ValueError(f"{nom}.{cle_liste}: exactement 4 ou 6 elements attendus, "
                             f"recu {len(items)}")
        for i, it in enumerate(items):
            _champs(it, champs, f"{nom}.{cle_liste}[{i}]")
    return data


# --- Contexte theme / zone -------------------------------------------------------
def _prs(slide):
    return slide.part.package.presentation_part.presentation


def _ctx(slide):
    prs = _prs(slide)
    tc = D.theme_colors(prs)
    manque = [k for k in ("dk1", "lt1", "lt2", "dk2", "accent1") if k not in tc]
    if manque:
        raise ValueError(f"theme du deck incomplet, couleurs absentes : {manque}")
    accents = [tc[k] for k in ("accent1", "accent2", "accent3", "accent4", "accent5",
                               "accent6") if k in tc]
    return {"ink": tc["dk1"], "paper": tc["lt1"], "panel": tc["lt2"], "muted": tc["dk2"],
            "accents": accents, "font": D.police_theme(prs)}


def _zone_defaut(slide):
    prs = _prs(slide)
    W, H = Emu(prs.slide_width).inches, Emu(prs.slide_height).inches
    haut = MARGE_H
    titre = slide.shapes.title
    if titre is not None and titre.top is not None and titre.height is not None:
        haut = Emu(titre.top + titre.height).inches + GAP
    return (MARGE_H, haut, W - 2 * MARGE_H, H - BAS_RESERVE - haut)


def _zone_in(slide, zone):
    if zone is None:
        return _zone_defaut(slide)
    prs = _prs(slide)
    x, y, w, h = (Emu(int(v)).inches for v in zone)
    W, H = Emu(prs.slide_width).inches, Emu(prs.slide_height).inches
    if w <= 0 or h <= 0 or x < 0 or y < 0 or x + w > W + 0.02 or y + h > H + 0.02:
        raise ValueError(f"zone {tuple(round(v, 2) for v in (x, y, w, h))} in invalide "
                         f"ou hors de la slide {W:.2f}x{H:.2f}")
    return (x, y, w, h)


# --- Mesure et pose du texte -----------------------------------------------------
def _h(texte, w, taille):
    """Hauteur estimee : meme modele que verifier_debordements_texte."""
    return D.estimer_lignes(texte, w, taille, cpi_ref=CPI_PESSIMISTE) * (taille * 0.017 + 4 / 72)


def _fit(textes, w, tmax, tmin, hauteur_fn, h_dispo, quoi):
    """Plus grande taille (pas 0,5 pt) telle que hauteur_fn(taille) <= h_dispo."""
    t, _ = D.ajuster_police(textes, w, tmax, tmin, lambda s, _n: hauteur_fn(s) <= h_dispo)
    if hauteur_fn(t) > h_dispo + 1e-9:
        raise ValueError(
            f"{quoi}: le texte ne tient pas dans le cadre meme a {t} pt "
            f"({hauteur_fn(t):.2f} in necessaires, {h_dispo:.2f} in disponibles) : "
            "raccourcir les textes ou repartir sur deux slides")
    return t


def _exiger(h, h_dispo, quoi):
    if h > h_dispo + 1e-9:
        raise ValueError(f"{quoi}: {h:.2f} in necessaires, {h_dispo:.2f} in disponibles "
                         "dans la zone : zone trop petite ou contenu trop long")


def _para(ctx, texte, taille, couleur, bold=False, align=PP_ALIGN.LEFT):
    return (texte, dict(size=taille, bold=bold, color=couleur, align=align,
                        space_after=4, font=ctx["font"]))


def _texte(slide, ctx, x, y, w, paras, anchor=MSO_ANCHOR.TOP):
    """Pose les paragraphes (texte, opts) dans une boite haute de leur estimation."""
    h = sum(_h(t, w, o["size"]) for t, o in paras) + 0.04
    D.add_text(slide, x, y, w, h, paras, anchor=anchor)
    return h


def _rect(slide, x, y, w, h, fill, rounded=False, radius=0.06):
    return D.add_rect(slide, x, y, w, h, fill=fill, rounded=rounded, radius=radius)


def _puces(textes):
    return ["• " + t for t in textes]


# --- Recettes --------------------------------------------------------------------
def kpi_row(slide, data, zone, ctx):
    x, y, w, h = zone
    kpis = data["kpis"]
    n = len(kpis)
    cw = (w - GAP * (n - 1)) / n
    inner = cw - 0.3
    t_val = _fit([k["value"] for k in kpis], inner, D.TYPE["kpi"], D.TYPE["h2"],
                 lambda s: max(_h(k["value"], inner, s) for k in kpis), 1.2, "kpi_row.value")
    h_val = max(_h(k["value"], inner, t_val) for k in kpis)

    def bas(s):
        return max(_h(k["label"], inner, s) + (_h(k["note"], inner, s) if k.get("note") else 0)
                   for k in kpis)

    t_txt = _fit([k["label"] for k in kpis], inner, T_MAX, T_MIN, bas,
                 max(h - h_val - 0.5, 0), "kpi_row.label")
    ch = h_val + bas(t_txt) + 0.5
    _exiger(ch, h, "kpi_row")
    top = y + (h - ch) / 2
    for i, k in enumerate(kpis):
        cx = x + i * (cw + GAP)
        acc = ctx["accents"][i % len(ctx["accents"])]
        _rect(slide, cx, top, cw, ch, ctx["panel"], rounded=True)
        _rect(slide, cx, top, cw, 0.07, acc)
        _texte(slide, ctx, cx + 0.15, top + 0.22, inner, [
            _para(ctx, k["value"], t_val, acc, bold=True, align=PP_ALIGN.CENTER)])
        paras = [_para(ctx, k["label"], t_txt, ctx["ink"], bold=True, align=PP_ALIGN.CENTER)]
        if k.get("note"):
            paras.append(_para(ctx, k["note"], t_txt, ctx["muted"], align=PP_ALIGN.CENTER))
        _texte(slide, ctx, cx + 0.15, top + 0.22 + h_val + 0.08, inner, paras)


def comparison(slide, data, zone, ctx):
    x, y, w, h = zone
    cols = data["columns"]
    n = len(cols)
    cw = (w - GAP * (n - 1)) / n
    inner = cw - 0.4
    t_tit = _fit([c["title"] for c in cols], inner, D.TYPE["h3"], D.TYPE["small"],
                 lambda s: max(_h(c["title"], inner, s) for c in cols), 0.8, "comparison.title")
    band = max(_h(c["title"], inner, t_tit) for c in cols) + 0.3

    def corps(s):
        return max(sum(_h(p, inner, s) for p in _puces(c["points"])) for c in cols)

    t_pts = _fit([p for c in cols for p in _puces(c["points"])], inner, T_MAX, T_MIN,
                 corps, h - band - 0.3, "comparison.points")
    for i, c in enumerate(cols):
        cx = x + i * (cw + GAP)
        acc = ctx["accents"][i % len(ctx["accents"])]
        _rect(slide, cx, y, cw, h, ctx["panel"], rounded=True)
        _rect(slide, cx, y, cw, band, acc)
        _texte(slide, ctx, cx + 0.2, y + 0.15, inner,
               [_para(ctx, c["title"], t_tit, ctx["paper"], bold=True)])
        _texte(slide, ctx, cx + 0.2, y + band + 0.15, inner,
               [_para(ctx, p, t_pts, ctx["ink"]) for p in _puces(c["points"])])


def timeline(slide, data, zone, ctx):
    x, y, w, h = zone
    steps = data["steps"]
    n = len(steps)
    cw = w / n
    inner = cw - 0.1

    def date_h(s):
        return max(_h(st["date"], inner, s) for st in steps)

    def bas_h(s):
        return max(_h(st["title"], inner, s) + (_h(st["text"], inner, s) if st.get("text") else 0)
                   for st in steps)

    def total(s):
        return date_h(s) + 0.1 + 0.2 + 0.1 + bas_h(s)

    t = _fit([st["title"] for st in steps] + [st["date"] for st in steps], inner,
             T_MAX, T_MIN, total, h, "timeline")
    top = y + (h - total(t)) / 2
    ligne_y = top + date_h(t) + 0.1 + 0.08
    _rect(slide, x + cw / 2, ligne_y, w - cw, 0.04, ctx["muted"])
    for i, st in enumerate(steps):
        cx = x + i * cw
        acc = ctx["accents"][i % len(ctx["accents"])]
        _texte(slide, ctx, cx + 0.05, top, inner,
               [_para(ctx, st["date"], t, acc, bold=True, align=PP_ALIGN.CENTER)])
        _rect(slide, cx + cw / 2 - 0.1, ligne_y - 0.08, 0.2, 0.2, acc, rounded=True, radius=0.5)
        paras = [_para(ctx, st["title"], t, ctx["ink"], bold=True, align=PP_ALIGN.CENTER)]
        if st.get("text"):
            paras.append(_para(ctx, st["text"], t, ctx["muted"], align=PP_ALIGN.CENTER))
        _texte(slide, ctx, cx + 0.05, ligne_y + 0.2 + 0.1, inner, paras)


def process(slide, data, zone, ctx):
    x, y, w, h = zone
    steps = data["steps"]
    n = len(steps)
    cw = w / n
    inner = cw - 0.2
    chev_h = 0.7

    def bas(s):
        return max(_h(st["title"], inner, s) + _h(st["text"], inner, s) for st in steps)

    t = _fit([st["title"] + " " + st["text"] for st in steps], inner, T_MAX, T_MIN, bas,
             h - chev_h - 0.2, "process")
    for i, st in enumerate(steps):
        cx = x + i * cw
        acc = ctx["accents"][i % len(ctx["accents"])]
        D.add_forme(slide, "chevron", cx, y, cw - 0.05, chev_h, fill=acc)
        D.add_text(slide, cx, y, cw - 0.05, chev_h,
                   [_para(ctx, str(i + 1), D.TYPE["h2"], ctx["paper"], bold=True,
                          align=PP_ALIGN.CENTER)], anchor=MSO_ANCHOR.MIDDLE)
        _texte(slide, ctx, cx + 0.1, y + chev_h + 0.2, inner, [
            _para(ctx, st["title"], t, ctx["ink"], bold=True),
            _para(ctx, st["text"], t, ctx["muted"])])


def _grille_cartes(slide, ctx, zone, items, cols, rows, quoi):
    x, y, w, h = zone
    cw = (w - GAP * (cols - 1)) / cols
    ch = (h - GAP * (rows - 1)) / rows
    inner = cw - 0.45

    def bloc(s):
        return max(_h(it["title"], inner, s + 2) + _h(it["text"], inner, s) for it in items)

    t = _fit([it["title"] + " " + it["text"] for it in items], inner, T_MAX, T_MIN, bloc,
             ch - 0.4, quoi)
    for i, it in enumerate(items):
        cx = x + (i % cols) * (cw + GAP)
        cy = y + (i // cols) * (ch + GAP)
        acc = ctx["accents"][i % len(ctx["accents"])]
        _rect(slide, cx, cy, cw, ch, ctx["panel"], rounded=True)
        _rect(slide, cx, cy, 0.07, ch, acc)
        _texte(slide, ctx, cx + 0.3, cy + 0.2, inner, [
            _para(ctx, it["title"], t + 2, acc, bold=True),
            _para(ctx, it["text"], t, ctx["ink"])])


def cards_grid(slide, data, zone, ctx):
    cartes = data["cards"]
    _grille_cartes(slide, ctx, zone, cartes, 2 if len(cartes) == 4 else 3, 2, "cards_grid")


def matrix_2x2(slide, data, zone, ctx):
    x, y, w, h = zone
    t_ax = D.TYPE["small"]
    h_ax = max(_h(data["y_axis"], w, t_ax), _h(data["x_axis"], w, t_ax)) + 0.04
    _texte(slide, ctx, x, y, w, [_para(ctx, "↑ " + data["y_axis"], t_ax, ctx["muted"],
                                       bold=True)])
    _exiger(h_ax * 2 + GAP * 2, h, "matrix_2x2")
    _grille_cartes(slide, ctx, (x, y + h_ax + 0.1, w, h - 2 * h_ax - 0.2), data["quadrants"],
                   2, 2, "matrix_2x2")
    _texte(slide, ctx, x, y + h - h_ax, w, [_para(ctx, data["x_axis"] + " →", t_ax,
                                                  ctx["muted"], bold=True,
                                                  align=PP_ALIGN.RIGHT)])


def quote(slide, data, zone, ctx):
    x, y, w, h = zone
    inner = w - 1.8
    t_q = _fit([data["quote"]], inner, D.TYPE["title"], T_MIN,
               lambda s: _h(data["quote"], inner, s), max(h - 1.9, 0.3), "quote")
    h_q = _h(data["quote"], inner, t_q)
    auteur = data["author"] + (" — " + data["role"] if data.get("role") else "")
    h_a = _h(auteur, inner, D.TYPE["h3"])
    ph = 0.45 + h_q + 0.25 + h_a + 0.45
    _exiger(ph, h, "quote")
    top = y + (h - ph) / 2
    _rect(slide, x, top, w, ph, ctx["muted"], rounded=True, radius=0.06)
    D.add_text(slide, x + 0.3, top + 0.1, 0.9, 0.85,
               [_para(ctx, "“", D.TYPE["kpi"], ctx["accents"][0], bold=True)])
    _texte(slide, ctx, x + 1.2, top + 0.45, inner, [_para(ctx, data["quote"], t_q, ctx["paper"],
                                                          bold=True)])
    _texte(slide, ctx, x + 1.2, top + 0.45 + h_q + 0.25, inner,
           [_para(ctx, auteur, D.TYPE["h3"], ctx["paper"])])


def highlight(slide, data, zone, ctx):
    x, y, w, h = zone
    lw = w * 0.36
    rw = w - lw - GAP
    inner_l = lw - 0.4
    t_v = _fit([data["value"]], inner_l, D.TYPE["kpi"], D.TYPE["h2"],
               lambda s: _h(data["value"], inner_l, s), 1.2, "highlight.value")
    h_v = _h(data["value"], inner_l, t_v)
    h_l = _h(data["label"], inner_l, D.TYPE["h3"])
    inner_r = rw - 0.6
    puces = _puces(data["points"])
    t_p = _fit(puces, inner_r, T_MAX, T_MIN,
               lambda s: sum(_h(p, inner_r, s) for p in puces), h - 0.5, "highlight.points")
    h_p = sum(_h(p, inner_r, t_p) for p in puces)
    ph = max(h_v + h_l + 0.5, h_p + 0.5)
    _exiger(ph, h, "highlight")
    top = y + (h - ph) / 2
    acc = ctx["accents"][0]
    _rect(slide, x, top, lw, ph, acc, rounded=True)
    _texte(slide, ctx, x + 0.2, top + 0.25, inner_l,
           [_para(ctx, data["value"], t_v, ctx["paper"], bold=True, align=PP_ALIGN.CENTER),
            _para(ctx, data["label"], D.TYPE["h3"], ctx["paper"], align=PP_ALIGN.CENTER)])
    _rect(slide, x + lw + GAP, top, rw, ph, ctx["panel"], rounded=True)
    _rect(slide, x + lw + GAP, top, 0.07, ph, acc)
    _texte(slide, ctx, x + lw + GAP + 0.3, top + 0.25, inner_r,
           [_para(ctx, p, t_p, ctx["ink"]) for p in puces])


RECETTES = {
    "kpi_row": (kpi_row, SPEC["kpi_row"]["doc"]),
    "comparison": (comparison, SPEC["comparison"]["doc"]),
    "timeline": (timeline, SPEC["timeline"]["doc"]),
    "process": (process, SPEC["process"]["doc"]),
    "cards_grid": (cards_grid, SPEC["cards_grid"]["doc"]),
    "quote": (quote, SPEC["quote"]["doc"]),
    "highlight": (highlight, SPEC["highlight"]["doc"]),
    "matrix_2x2": (matrix_2x2, SPEC["matrix_2x2"]["doc"]),
}


def lister():
    """Une ligne par recette : nom + schema de data (a lire a la place du catalogue)."""
    return "\n".join(f"{nom}: {doc}" for nom, (_f, doc) in RECETTES.items())


# --- Controle final --------------------------------------------------------------
def _controler(slide, zone, avant):
    prs = _prs(slide)
    num = [s.slide_id for s in prs.slides].index(slide.slide_id) + 1
    marque = re.compile(rf"\s*slide {num}:")
    defauts = [p for p in list(D.verifier_geometrie(prs)) + list(D.verifier_debordements_texte(prs))
               if marque.match(str(p))]
    zx, zy, zw, zh = zone
    tol = 0.03
    for sh in slide.shapes:
        if sh.shape_id in avant:
            continue
        l, t = Emu(sh.left).inches, Emu(sh.top).inches
        r, b = l + Emu(sh.width).inches, t + Emu(sh.height).inches
        if l < zx - tol or t < zy - tol or r > zx + zw + tol or b > zy + zh + tol:
            defauts.append(f"slide {num}: '{sh.name}' sort de la zone de la recette "
                           f"(l={l:.2f} t={t:.2f} r={r:.2f} b={b:.2f})")
    if defauts:
        raise RecetteDefaut("recette refusee, defauts : " + " | ".join(defauts))


def recette(slide, nom, data, zone=None):
    """Dessine la recette `nom` sur `slide` depuis `data` seule. Voir lister()."""
    if nom not in RECETTES:
        raise ValueError(f"recette inconnue {nom!r} ; disponibles : {sorted(RECETTES)}")
    data = copy.deepcopy(valider(nom, data))
    zone_in = _zone_in(slide, zone)
    ctx = _ctx(slide)
    avant = {sh.shape_id for sh in slide.shapes}
    try:
        RECETTES[nom][0](slide, data, zone_in, ctx)
        _controler(slide, zone_in, avant)
    except Exception:
        for sh in list(slide.shapes):  # atomique : un refus ne laisse aucune forme
            if sh.shape_id not in avant:
                sh._element.getparent().remove(sh._element)
        raise
